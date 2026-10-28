from __future__ import annotations

import argparse
import json
import re
import time
import traceback
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from olbench.preprocess import chunk_text
from olbench.extractors import (
    yake_extract,
    tfidf,
    noun_phrases,
    textrank,
    keybert_extract,
)


def normalize_term(value: object) -> str:
    text = str(value or "").casefold().strip()
    text = text.replace("‐", "-").replace("‑", "-").replace("–", "-").replace("—", "-")
    text = re.sub(r"(?<=\w)-(?=\w)", " ", text)
    text = re.sub(r"[^a-z0-9&/+ ]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def load_or_create_chunks(cleaned_file: Path, chunks_file: Path, max_words: int, overlap: int) -> pd.DataFrame:
    if chunks_file.exists():
        chunks = pd.read_csv(chunks_file)
        required = {"chunk_id", "text"}
        if not required.issubset(chunks.columns):
            raise ValueError(f"{chunks_file} must contain: {sorted(required)}")
        return chunks

    text = cleaned_file.read_text(encoding="utf-8")
    chunks = chunk_text(text, max_words=max_words, overlap_sentences=overlap)
    chunks_file.parent.mkdir(parents=True, exist_ok=True)
    chunks.to_csv(chunks_file, index=False)
    return chunks


def load_ground_truth(workbook: Path):
    gt = pd.read_excel(workbook, sheet_name="S1_GT_TERMS", engine="openpyxl")
    negatives = pd.read_excel(workbook, sheet_name="S7_NEGATIVE_TERMS", engine="openpyxl")

    gt = gt[gt["keep_for_tasks"].astype(str).str.contains("term_extraction", case=False, na=False)].copy()

    # Strict benchmark excludes rows explicitly marked for manual review.
    strict_gt = gt[gt["grounded_status_v2"].astype(str).str.casefold() != "needs_manual_review"].copy()
    review_gt = gt[gt["grounded_status_v2"].astype(str).str.casefold() == "needs_manual_review"].copy()

    for frame, col in [(strict_gt, "term_canonical"), (review_gt, "term_canonical")]:
        frame["normalized_gt"] = frame[col].map(normalize_term)

    negatives["normalized_negative"] = negatives["candidate"].map(normalize_term)
    return strict_gt, review_gt, negatives


def aggregate_candidates(raw: pd.DataFrame, method: str) -> pd.DataFrame:
    if raw.empty:
        return pd.DataFrame(columns=[
            "method", "candidate", "normalized_candidate", "score", "global_rank",
            "occurrence_count", "chunk_count", "best_chunk_id", "source_sentence", "context_text"
        ])

    work = raw.copy()
    work["normalized_candidate"] = work["raw_term"].map(normalize_term)
    work = work[work["normalized_candidate"].str.len() > 0].copy()
    work["score"] = pd.to_numeric(work["score"], errors="coerce").fillna(0.0)

    # Aggregate duplicate candidates across chunks. Maximum score is used because
    # all framework extractors expose larger values as stronger candidates.
    work = work.sort_values(["score", "rank"], ascending=[False, True])
    best = work.drop_duplicates("normalized_candidate", keep="first").set_index("normalized_candidate")
    stats = work.groupby("normalized_candidate").agg(
        occurrence_count=("normalized_candidate", "size"),
        chunk_count=("chunk_id", "nunique"),
        aggregate_score=("score", "max"),
    )

    result = stats.join(best[["raw_term", "chunk_id", "source_sentence", "context_text"]])
    result = result.reset_index().rename(columns={
        "raw_term": "candidate",
        "aggregate_score": "score",
        "chunk_id": "best_chunk_id",
    })
    result.insert(0, "method", method)
    result = result.sort_values(
        ["score", "occurrence_count", "normalized_candidate"],
        ascending=[False, False, True],
    ).reset_index(drop=True)
    result["global_rank"] = range(1, len(result) + 1)
    return result[[
        "method", "candidate", "normalized_candidate", "score", "global_rank",
        "occurrence_count", "chunk_count", "best_chunk_id", "source_sentence", "context_text"
    ]]


def evaluate_method(candidates: pd.DataFrame, strict_gt: pd.DataFrame,
                    review_gt: pd.DataFrame, negatives: pd.DataFrame):
    gt_lookup = strict_gt.drop_duplicates("normalized_gt").set_index("normalized_gt")
    review_lookup = review_gt.drop_duplicates("normalized_gt").set_index("normalized_gt")
    negative_lookup = negatives.drop_duplicates("normalized_negative").set_index("normalized_negative")

    details = candidates.copy()
    details["matched_gt_id"] = details["normalized_candidate"].map(
        lambda x: gt_lookup.loc[x, "gt_id"] if x in gt_lookup.index else ""
    )
    details["match_type"] = details["normalized_candidate"].map(
        lambda x: "strict_exact" if x in gt_lookup.index else
                  "manual_review_exact" if x in review_lookup.index else "unmatched"
    )
    details["evaluation_flag"] = details["match_type"].map(
        {"strict_exact": "TP", "manual_review_exact": "REVIEW", "unmatched": "FP"}
    )
    details["gt_category"] = details["normalized_candidate"].map(
        lambda x: gt_lookup.loc[x, "gt_category"] if x in gt_lookup.index else ""
    )
    details["negative_benchmark_hit"] = details["normalized_candidate"].isin(negative_lookup.index)

    predicted = set(details["normalized_candidate"])
    gold = set(strict_gt["normalized_gt"])
    tp = len(predicted & gold)
    fp = len(predicted - gold)
    fn = len(gold - predicted)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0

    metric = {
        "method": details["method"].iloc[0] if len(details) else "unknown",
        "candidate_count": len(details),
        "strict_gold_count": len(gold),
        "manual_review_gold_count": review_gt["normalized_gt"].nunique(),
        "true_positives": tp,
        "false_positives": fp,
        "false_negatives": fn,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "manual_review_matches": int((details["match_type"] == "manual_review_exact").sum()),
        "negative_term_hits": int(details["negative_benchmark_hit"].sum()),
    }

    for k in (25, 50, 100):
        top = details.nsmallest(k, "global_rank")
        top_tp = int((top["match_type"] == "strict_exact").sum())
        metric[f"precision_at_{k}"] = top_tp / len(top) if len(top) else 0.0
        metric[f"recall_at_{k}"] = top_tp / len(gold) if gold else 0.0

    # Since all current S1 gold priorities are high, priority-weighted recall
    # equals strict recall. This calculation remains generic for future weights.
    weights = {"high": 3.0, "medium": 2.0, "low": 1.0}
    gt_weight = strict_gt["priority"].astype(str).str.casefold().map(weights).fillna(1.0)
    matched_mask = strict_gt["normalized_gt"].isin(predicted)
    metric["priority_weighted_recall"] = (
        float(gt_weight[matched_mask].sum() / gt_weight.sum()) if gt_weight.sum() else 0.0
    )

    fn_rows = strict_gt[~strict_gt["normalized_gt"].isin(predicted)].copy()
    fn_report = fn_rows[["gt_id", "term_canonical", "normalized_gt", "gt_category", "priority",
                         "grounded_status_v2", "evidence_chapter", "evidence_page_pdf"]]
    fn_report.insert(0, "method", metric["method"])
    fn_report["evaluation_flag"] = "FN"

    categories = []
    for category, group in strict_gt.groupby("gt_category"):
        category_gold = set(group["normalized_gt"])
        category_tp = len(category_gold & predicted)
        categories.append({
            "method": metric["method"],
            "gt_category": category,
            "gold_count": len(category_gold),
            "matched_count": category_tp,
            "recall": category_tp / len(category_gold) if category_gold else 0.0,
        })

    negative_hits = details[details["negative_benchmark_hit"]].copy()
    return metric, details, fn_report, pd.DataFrame(categories), negative_hits


def style_workbook(path: Path):
    wb = load_workbook(path)
    navy = "1F4E78"
    teal = "DDEBF7"
    green = "008000"

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30

        for col_idx, column in enumerate(ws.iter_cols(1, ws.max_column), 1):
            values = [str(c.value) if c.value is not None else "" for c in list(column)[:200]]
            width = min(max(max((len(v) for v in values), default=8) + 2, 11), 48)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
                if ws.title != "README":
                    cell.font = Font(color=green)

        if ws.title == "Metrics Summary":
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(color="000000")
            for col in range(8, ws.max_column + 1):
                for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                    for c in cell:
                        if isinstance(c.value, float):
                            c.number_format = "0.0%"

    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="Run and evaluate non-LLM term extraction methods.")
    parser.add_argument("--cleaned", required=True, type=Path)
    parser.add_argument("--chunks", required=True, type=Path)
    parser.add_argument("--ground-truth", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--top-k-per-chunk", type=int, default=30)
    parser.add_argument("--max-words", type=int, default=220)
    parser.add_argument("--overlap-sentences", type=int, default=1)
    args = parser.parse_args()

    args.out.mkdir(parents=True, exist_ok=True)
    chunks = load_or_create_chunks(args.cleaned, args.chunks, args.max_words, args.overlap_sentences)
    strict_gt, review_gt, negatives = load_ground_truth(args.ground_truth)

    methods = [
        ("yake", yake_extract, {"top_k": args.top_k_per_chunk}),
        ("tfidf", tfidf, {"top_k": args.top_k_per_chunk}),
        ("noun_phrases", noun_phrases, {"top_k": args.top_k_per_chunk}),
        ("textrank", textrank, {"top_k": args.top_k_per_chunk}),
        ("keybert_minilm", keybert_extract, {
            "top_k": args.top_k_per_chunk,
            "model": "all-MiniLM-L6-v2",
        }),
        ("keybert_mpnet", keybert_extract, {
            "top_k": args.top_k_per_chunk,
            "model": "all-mpnet-base-v2",
        }),
    ]

    metrics, all_unique, all_details, all_fn, all_categories, all_negative_hits = [], [], [], [], [], []
    manifest = {
        "cleaned_input": str(args.cleaned.resolve()),
        "chunks_input": str(args.chunks.resolve()),
        "ground_truth": str(args.ground_truth.resolve()),
        "strict_gt_rule": "S1 term_extraction rows excluding needs_manual_review",
        "matching_rule": "case-insensitive normalized exact match",
        "chunk_count": len(chunks),
        "methods": [],
    }

    for method, function, params in methods:
        method_dir = args.out / "by_method" / method
        method_dir.mkdir(parents=True, exist_ok=True)
        started = time.time()
        print(f"\nRunning {method} ...")
        try:
            raw = function(chunks, **params)
            unique = aggregate_candidates(raw, method)
            metric, details, fn_report, category_report, negative_hits = evaluate_method(
                unique, strict_gt, review_gt, negatives
            )

            raw.to_csv(method_dir / "candidates_all_occurrences.csv", index=False)
            unique.to_csv(method_dir / "candidates_unique_ranked.csv", index=False)
            details.to_csv(method_dir / "evaluation_predictions.csv", index=False)
            fn_report.to_csv(method_dir / "false_negatives.csv", index=False)
            category_report.to_csv(method_dir / "category_recall.csv", index=False)
            negative_hits.to_csv(method_dir / "negative_term_hits.csv", index=False)
            pd.DataFrame([metric]).to_csv(method_dir / "metrics.csv", index=False)

            metrics.append(metric)
            all_unique.append(unique)
            all_details.append(details)
            all_fn.append(fn_report)
            all_categories.append(category_report)
            all_negative_hits.append(negative_hits)
            manifest["methods"].append({
                "method": method,
                "status": "completed",
                "parameters": params,
                "duration_seconds": round(time.time() - started, 3),
                "unique_candidates": len(unique),
            })
            print(f"Completed {method}: {len(unique)} unique candidates, F1={metric['f1']:.4f}")
        except Exception as exc:
            error_text = traceback.format_exc()
            (method_dir / "ERROR.txt").write_text(error_text, encoding="utf-8")
            manifest["methods"].append({
                "method": method,
                "status": "failed",
                "parameters": params,
                "duration_seconds": round(time.time() - started, 3),
                "error": str(exc),
            })
            print(f"FAILED {method}: {exc}")

    metrics_df = pd.DataFrame(metrics).sort_values("f1", ascending=False) if metrics else pd.DataFrame()
    candidates_df = pd.concat(all_unique, ignore_index=True) if all_unique else pd.DataFrame()
    details_df = pd.concat(all_details, ignore_index=True) if all_details else pd.DataFrame()
    fn_df = pd.concat(all_fn, ignore_index=True) if all_fn else pd.DataFrame()
    category_df = pd.concat(all_categories, ignore_index=True) if all_categories else pd.DataFrame()
    negative_df = pd.concat(all_negative_hits, ignore_index=True) if all_negative_hits else pd.DataFrame()

    metrics_df.to_csv(args.out / "metrics_summary.csv", index=False)
    candidates_df.to_csv(args.out / "all_methods_candidates.csv", index=False)
    details_df.to_csv(args.out / "all_methods_evaluation.csv", index=False)
    fn_df.to_csv(args.out / "all_methods_false_negatives.csv", index=False)
    category_df.to_csv(args.out / "all_methods_category_recall.csv", index=False)
    negative_df.to_csv(args.out / "all_methods_negative_hits.csv", index=False)

    workbook = args.out / "non_llm_extraction_report.xlsx"
    readme = pd.DataFrame({
        "Field": [
            "Purpose", "Strict GT count", "Manual-review GT count", "Matching rule",
            "Precision meaning", "Recall meaning", "Important limitation"
        ],
        "Value": [
            "Compare non-LLM extraction methods against S1_GT_TERMS",
            strict_gt["normalized_gt"].nunique(),
            review_gt["normalized_gt"].nunique(),
            "Case-insensitive normalized exact match",
            "Exact GT matches divided by all unique extracted candidates",
            "Exact GT matches divided by strict grounded GT terms",
            "Ground truth is a silver benchmark and may not enumerate every valid domain term; unmatched candidates require review.",
        ],
    })

    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        metrics_df.to_excel(writer, sheet_name="Metrics Summary", index=False)
        category_df.to_excel(writer, sheet_name="Category Recall", index=False)
        details_df.to_excel(writer, sheet_name="Prediction Evaluation", index=False)
        fn_df.to_excel(writer, sheet_name="False Negatives", index=False)
        negative_df.to_excel(writer, sheet_name="Negative Hits", index=False)
        candidates_df.to_excel(writer, sheet_name="All Candidates", index=False)

    style_workbook(workbook)
    (args.out / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print("\nFinished.")
    print(f"Metrics: {args.out / 'metrics_summary.csv'}")
    print(f"Excel report: {workbook}")
    print(f"Manifest: {args.out / 'run_manifest.json'}")


if __name__ == "__main__":
    main()
