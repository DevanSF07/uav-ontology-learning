from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd


PROJECT_DIR = Path("/home/faujdar/onto_gen/ontology_experiment_framework")
DEFAULT_CANDIDATES = PROJECT_DIR / "stage_outputs/04_gliner_filtering/best_gliner_retained.csv"
DEFAULT_GT = PROJECT_DIR / "data/Ground_Truth.xlsx"
DEFAULT_MATCHING_DIR = PROJECT_DIR / "stage_outputs/05_clustering_filtering/cluster_matching"
DEFAULT_OUTPUT = PROJECT_DIR / "stage_outputs/05_clustering_filtering/cluster_removal_evaluation"

GT_SHEET = "S1_GT_TERMS"
NEGATIVE_SHEET = "S7_NEGATIVE_TERMS"
MANUAL_REVIEW_STATUS = "needs_manual_review"
PRIORITY_WEIGHTS = {"high": 3.0, "medium": 2.0, "low": 1.0}


def normalize_term(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    text = str(value).casefold().strip()
    for dash in ("‐", "‑", "‒", "–", "—"):
        text = text.replace(dash, "-")
    text = re.sub(r"(?<=\w)-(?=\w)", " ", text)
    text = re.sub(r"[^a-z0-9&/+ ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def read_csv(path: Path, required: bool = True) -> pd.DataFrame:
    if not path.is_file():
        if required:
            raise FileNotFoundError(path)
        return pd.DataFrame()
    return pd.read_csv(path, encoding="utf-8")


def validate_columns(frame: pd.DataFrame, required: set[str], label: str) -> None:
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"{label} is missing columns: {sorted(missing)}")


def load_ground_truth(gt_file: Path):
    if not gt_file.is_file():
        raise FileNotFoundError(gt_file)

    gt = pd.read_excel(gt_file, sheet_name=GT_SHEET, engine="openpyxl")
    negatives = pd.read_excel(gt_file, sheet_name=NEGATIVE_SHEET, engine="openpyxl")

    gt = gt[
        gt["keep_for_tasks"].astype(str).str.contains(
            "term_extraction", case=False, na=False
        )
    ].copy()
    gt["normalized_gt"] = gt["term_canonical"].map(normalize_term)
    gt = gt[gt["normalized_gt"].ne("")]

    status = gt["grounded_status_v2"].fillna("").astype(str).str.casefold()
    strict_gt = gt[status.ne(MANUAL_REVIEW_STATUS)].drop_duplicates("normalized_gt").copy()
    review_gt = gt[status.eq(MANUAL_REVIEW_STATUS)].drop_duplicates("normalized_gt").copy()

    negatives["normalized_negative"] = negatives["candidate"].map(normalize_term)
    negatives = negatives[negatives["normalized_negative"].ne("")]
    return strict_gt, review_gt, negatives


def prepare_candidates(
    candidates_file: Path,
    strict_gt: pd.DataFrame,
    review_gt: pd.DataFrame,
    negatives: pd.DataFrame,
) -> pd.DataFrame:
    candidates = read_csv(candidates_file)
    validate_columns(candidates, {"candidate", "normalized_candidate"}, "candidate inventory")

    candidates = candidates.copy()
    candidates["normalized_candidate"] = candidates["normalized_candidate"].map(normalize_term)
    candidates = candidates[candidates["normalized_candidate"].ne("")]
    candidates = candidates.drop_duplicates("normalized_candidate", keep="first").reset_index(drop=True)
    if "candidate_row_id" not in candidates.columns:
        candidates.insert(0, "candidate_row_id", np.arange(len(candidates), dtype=int))

    strict_lookup = strict_gt.set_index("normalized_gt")
    review_lookup = review_gt.set_index("normalized_gt")
    negative_set = set(negatives["normalized_negative"])

    candidates["match_type"] = candidates["normalized_candidate"].map(
        lambda term: "strict_exact" if term in strict_lookup.index else (
            "manual_review_exact" if term in review_lookup.index else "unmatched"
        )
    )
    candidates["evaluation_flag_before_filter"] = candidates["match_type"].map(
        {"strict_exact": "TP", "manual_review_exact": "REVIEW", "unmatched": "FP"}
    )
    candidates["matched_gt_id"] = candidates["normalized_candidate"].map(
        lambda term: strict_lookup.at[term, "gt_id"] if term in strict_lookup.index else (
            review_lookup.at[term, "gt_id"] if term in review_lookup.index else ""
        )
    )
    candidates["gt_category"] = candidates["normalized_candidate"].map(
        lambda term: strict_lookup.at[term, "gt_category"] if term in strict_lookup.index else (
            review_lookup.at[term, "gt_category"] if term in review_lookup.index else ""
        )
    )
    candidates["negative_benchmark_hit"] = candidates["normalized_candidate"].isin(negative_set)
    return candidates


def normalized_terms(frame: pd.DataFrame) -> set[str]:
    if frame.empty:
        return set()
    if "normalized_candidate" in frame.columns:
        return set(frame["normalized_candidate"].map(normalize_term)) - {""}
    if "candidate" in frame.columns:
        return set(frame["candidate"].map(normalize_term)) - {""}
    raise ValueError("Removal file must contain candidate or normalized_candidate")


def evaluate_inventory(
    experiment: str,
    retained: pd.DataFrame,
    removed: pd.DataFrame,
    strict_gt: pd.DataFrame,
    removal_source: str,
) -> dict:
    strict_unique = strict_gt.drop_duplicates("normalized_gt")
    gold = set(strict_unique["normalized_gt"])

    scored = retained[retained["match_type"].ne("manual_review_exact")]
    predicted = set(scored["normalized_candidate"])

    tp = len(predicted & gold)
    fp = len(predicted - gold)
    fn = len(gold - predicted)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0

    weights = (
        strict_unique["priority"].fillna("low").astype(str).str.casefold()
        .map(PRIORITY_WEIGHTS).fillna(1.0)
    )
    weighted_recall = float(
        weights[strict_unique["normalized_gt"].isin(predicted)].sum() / weights.sum()
    ) if weights.sum() else 0.0

    tp_removed = int(removed["match_type"].eq("strict_exact").sum())
    fp_removed = int(removed["match_type"].eq("unmatched").sum())

    return {
        "experiment": experiment,
        "removal_source": removal_source,
        "input_candidate_count": len(retained) + len(removed),
        "retained_candidate_count": len(retained),
        "removed_candidate_count": len(removed),
        "true_positives": tp,
        "false_positives": fp,
        "false_negatives": fn,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "priority_weighted_recall": weighted_recall,
        "tp_removed": tp_removed,
        "fp_removed": fp_removed,
        "manual_review_removed": int(removed["match_type"].eq("manual_review_exact").sum()),
        "negative_hits_retained": int(retained["negative_benchmark_hit"].sum()),
        "negative_hits_removed": int(removed["negative_benchmark_hit"].sum()),
        "fp_removed_per_tp_removed": fp_removed / tp_removed if tp_removed else np.inf,
        "smoothed_filter_efficiency": fp_removed / (tp_removed + 1),
    }


def apply_removal(
    candidates: pd.DataFrame,
    terms_to_remove: set[str],
    experiment: str,
    strict_gt: pd.DataFrame,
    removal_source: str,
) -> tuple[dict, pd.DataFrame, pd.DataFrame]:
    remove_mask = candidates["normalized_candidate"].isin(terms_to_remove)
    removed = candidates[remove_mask].copy()
    retained = candidates[~remove_mask].copy()

    removed["removal_experiment"] = experiment
    removed["removal_source"] = removal_source
    removed["filter_decision"] = "removed"
    retained["removal_experiment"] = experiment
    retained["removal_source"] = removal_source
    retained["filter_decision"] = "retained"

    metrics = evaluate_inventory(
        experiment, retained, removed, strict_gt, removal_source
    )
    return metrics, retained, removed


def cluster_term_sets(consensus_members: pd.DataFrame) -> dict[str, set[str]]:
    if consensus_members.empty:
        return {}
    validate_columns(
        consensus_members,
        {"matched_cluster_id", "candidate"},
        "three_way_consensus_members.csv",
    )
    result = {}
    for cluster_id, group in consensus_members.groupby("matched_cluster_id", sort=True):
        result[str(cluster_id)] = normalized_terms(group)
    return result


def create_category_loss_summary(removed: pd.DataFrame, experiment: str) -> pd.DataFrame:
    strict_removed = removed[removed["match_type"].eq("strict_exact")].copy()
    if strict_removed.empty:
        return pd.DataFrame(columns=["experiment", "gt_category", "strict_terms_removed"])
    summary = (
        strict_removed.groupby("gt_category", dropna=False)
        .size()
        .rename("strict_terms_removed")
        .reset_index()
    )
    summary.insert(0, "experiment", experiment)
    return summary


def format_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        if worksheet.max_row and worksheet.max_column:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for column in range(1, worksheet.max_column + 1):
            lengths = [
                len(str(worksheet.cell(row, column).value or ""))
                for row in range(1, min(worksheet.max_row, 150) + 1)
            ]
            worksheet.column_dimensions[get_column_letter(column)].width = min(
                max(max(lengths, default=10) + 2, 11), 45
            )
        headers = {worksheet.cell(1, col).value: col for col in range(1, worksheet.max_column + 1)}
        for name in ("precision", "recall", "f1", "priority_weighted_recall"):
            if name in headers:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row, headers[name]).number_format = "0.00%"
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Remove terms returned by three-way equivalent cluster matching and/or "
            "HDBSCAN native noise, then recalculate GT metrics."
        )
    )
    parser.add_argument("--candidates", type=Path, default=DEFAULT_CANDIDATES)
    parser.add_argument("--ground-truth", type=Path, default=DEFAULT_GT)
    parser.add_argument("--matching-dir", type=Path, default=DEFAULT_MATCHING_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--mode",
        choices=("all", "consensus", "noise", "combined", "cluster_ablation"),
        default="all",
        help=(
            "all evaluates every strategy; consensus removes all three-way consensus terms; "
            "noise removes HDBSCAN label -1; combined removes both; cluster_ablation removes "
            "one matched consensus cluster at a time."
        ),
    )
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    retained_dir = args.output_dir / "retained"
    removed_dir = args.output_dir / "removed"
    retained_dir.mkdir(parents=True, exist_ok=True)
    removed_dir.mkdir(parents=True, exist_ok=True)

    strict_gt, review_gt, negatives = load_ground_truth(args.ground_truth)
    candidates = prepare_candidates(args.candidates, strict_gt, review_gt, negatives)

    consensus_file = args.matching_dir / "three_way_consensus_members.csv"
    noise_file = args.matching_dir / "hdbscan_noise_cluster.csv"
    consensus_members = read_csv(consensus_file, required=False)
    hdbscan_noise = read_csv(noise_file, required=False)

    consensus_terms = normalized_terms(consensus_members)
    noise_terms = normalized_terms(hdbscan_noise)
    combined_terms = consensus_terms | noise_terms
    per_cluster = cluster_term_sets(consensus_members)

    metrics_rows = []
    output_map: dict[str, tuple[pd.DataFrame, pd.DataFrame]] = {}
    category_summaries = []

    # Always evaluate the baseline.
    baseline_metrics, baseline_retained, baseline_removed = apply_removal(
        candidates,
        set(),
        "baseline_no_additional_removal",
        strict_gt,
        "none",
    )
    metrics_rows.append(baseline_metrics)
    output_map["baseline_no_additional_removal"] = (baseline_retained, baseline_removed)

    strategies: list[tuple[str, set[str], str]] = []
    if args.mode in {"all", "consensus"}:
        strategies.append((
            "remove_all_three_way_consensus_terms",
            consensus_terms,
            "three_way_consensus_members",
        ))
    if args.mode in {"all", "noise"}:
        strategies.append((
            "remove_hdbscan_native_noise",
            noise_terms,
            "hdbscan_cluster_minus_1",
        ))
    if args.mode in {"all", "combined"}:
        strategies.append((
            "remove_consensus_and_hdbscan_noise",
            combined_terms,
            "three_way_consensus_union_hdbscan_noise",
        ))
    if args.mode in {"all", "cluster_ablation"}:
        for cluster_id, terms in per_cluster.items():
            strategies.append((
                f"remove_consensus_cluster_{cluster_id}",
                terms,
                f"three_way_consensus_cluster:{cluster_id}",
            ))

    for experiment, terms, source in strategies:
        metric, retained, removed = apply_removal(
            candidates, terms, experiment, strict_gt, source
        )
        metrics_rows.append(metric)
        output_map[experiment] = (retained, removed)
        category_summaries.append(create_category_loss_summary(removed, experiment))

        retained.to_csv(retained_dir / f"{experiment}.csv", index=False, encoding="utf-8")
        removed.to_csv(removed_dir / f"{experiment}.csv", index=False, encoding="utf-8")

    metrics = pd.DataFrame(metrics_rows)
    baseline_f1 = float(metrics.loc[
        metrics["experiment"].eq("baseline_no_additional_removal"), "f1"
    ].iloc[0])
    baseline_precision = float(metrics.loc[
        metrics["experiment"].eq("baseline_no_additional_removal"), "precision"
    ].iloc[0])
    baseline_recall = float(metrics.loc[
        metrics["experiment"].eq("baseline_no_additional_removal"), "recall"
    ].iloc[0])

    metrics["f1_change"] = metrics["f1"] - baseline_f1
    metrics["precision_change"] = metrics["precision"] - baseline_precision
    metrics["recall_change"] = metrics["recall"] - baseline_recall
    metrics = metrics.sort_values(
        ["f1", "precision", "recall"], ascending=[False, False, False]
    ).reset_index(drop=True)
    metrics["rank_by_f1"] = np.arange(1, len(metrics) + 1)

    metrics.to_csv(args.output_dir / "cluster_removal_metrics.csv", index=False, encoding="utf-8")

    category_summary = (
        pd.concat(category_summaries, ignore_index=True)
        if category_summaries
        else pd.DataFrame(columns=["experiment", "gt_category", "strict_terms_removed"])
    )
    category_summary.to_csv(
        args.output_dir / "strict_gt_category_losses.csv", index=False, encoding="utf-8"
    )

    best_experiment = str(metrics.iloc[0]["experiment"])
    best_retained, best_removed = output_map[best_experiment]
    best_retained.to_csv(args.output_dir / "best_cluster_removal_retained.csv", index=False)
    best_removed.to_csv(args.output_dir / "best_cluster_removal_removed.csv", index=False)

    source_summary = pd.DataFrame([
        {
            "removal_set": "three_way_consensus_members",
            "unique_terms": len(consensus_terms),
            "terms_present_in_candidate_inventory": int(
                candidates["normalized_candidate"].isin(consensus_terms).sum()
            ),
        },
        {
            "removal_set": "hdbscan_native_noise",
            "unique_terms": len(noise_terms),
            "terms_present_in_candidate_inventory": int(
                candidates["normalized_candidate"].isin(noise_terms).sum()
            ),
        },
        {
            "removal_set": "union",
            "unique_terms": len(combined_terms),
            "terms_present_in_candidate_inventory": int(
                candidates["normalized_candidate"].isin(combined_terms).sum()
            ),
        },
        {
            "removal_set": "intersection_consensus_and_noise",
            "unique_terms": len(consensus_terms & noise_terms),
            "terms_present_in_candidate_inventory": int(
                candidates["normalized_candidate"].isin(consensus_terms & noise_terms).sum()
            ),
        },
    ])
    source_summary.to_csv(args.output_dir / "removal_set_summary.csv", index=False)

    report_path = args.output_dir / "cluster_removal_evaluation.xlsx"
    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        metrics.to_excel(writer, sheet_name="Metrics", index=False)
        source_summary.to_excel(writer, sheet_name="Removal Sets", index=False)
        category_summary.to_excel(writer, sheet_name="GT Category Losses", index=False)
        best_retained.to_excel(writer, sheet_name="Best Retained", index=False)
        best_removed.to_excel(writer, sheet_name="Best Removed", index=False)
    format_workbook(report_path)

    manifest = {
        "candidate_inventory": str(args.candidates),
        "ground_truth": str(args.ground_truth),
        "matching_directory": str(args.matching_dir),
        "mode": args.mode,
        "input_candidate_count": len(candidates),
        "three_way_consensus_term_count": len(consensus_terms),
        "hdbscan_native_noise_term_count": len(noise_terms),
        "combined_removal_term_count": len(combined_terms),
        "matched_consensus_cluster_count": len(per_cluster),
        "evaluated_experiment_count": len(metrics),
        "best_experiment": best_experiment,
        "best_f1": float(metrics.iloc[0]["f1"]),
        "baseline_f1": baseline_f1,
    }
    (args.output_dir / "cluster_removal_manifest.json").write_text(
        json.dumps(manifest, indent=2), encoding="utf-8"
    )

    display = [
        "rank_by_f1",
        "experiment",
        "retained_candidate_count",
        "removed_candidate_count",
        "true_positives",
        "false_positives",
        "precision",
        "recall",
        "f1",
        "f1_change",
        "tp_removed",
        "fp_removed",
        "smoothed_filter_efficiency",
    ]
    print("=" * 100)
    print("Cluster-term removal evaluation completed")
    print("=" * 100)
    print(metrics[display].to_string(index=False))
    print(f"\nBest experiment: {best_experiment}")
    print(f"Metrics: {args.output_dir / 'cluster_removal_metrics.csv'}")
    print(f"Best removed terms: {args.output_dir / 'best_cluster_removal_removed.csv'}")
    print(f"Report: {report_path}")


if __name__ == "__main__":
    main()
