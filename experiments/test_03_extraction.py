from __future__ import annotations

import json
import re
import time
import traceback
from pathlib import Path

import pandas as pd

from olbench.extractors import (
    keybert_extract,
    noun_phrases,
    textrank,
    tfidf,
    yake_extract,
)


# ---------------------------------------------------------------------------
# Project paths
# ---------------------------------------------------------------------------

PROJECT_DIR = Path(
    "ontology_experiment_framework"
)

CHUNKS_FILE = PROJECT_DIR / "stage_outputs/02_chunking/chunks.csv"
GROUND_TRUTH_FILE = PROJECT_DIR / "data/Ground_Truth.xlsx"
OUTPUT_DIR = PROJECT_DIR / "stage_outputs/03_non_llm_extraction"


# ---------------------------------------------------------------------------
# Experiment configuration
# ---------------------------------------------------------------------------

CANDIDATE_SELECTION_RULE = (
    "Retain all candidates returned by each extraction method. "
    "No user-defined top-k cutoff is applied."
)
GROUND_TRUTH_SHEET = "S1_GT_TERMS"
NEGATIVE_TERMS_SHEET = "S7_NEGATIVE_TERMS"
MANUAL_REVIEW_STATUS = "needs_manual_review"

PRIORITY_WEIGHTS = {
    "high": 3.0,
    "medium": 2.0,
    "low": 1.0,
}


# ---------------------------------------------------------------------------
# Text normalization
# ---------------------------------------------------------------------------


def normalize_term(value: object) -> str:
    """Normalize an extracted or ground-truth term for exact matching."""

    if value is None or pd.isna(value):
        return ""

    text = str(value).casefold().strip()

    text = (
        text.replace("‐", "-")
        .replace("‑", "-")
        .replace("‒", "-")
        .replace("–", "-")
        .replace("—", "-")
    )

    text = re.sub(r"(?<=\w)-(?=\w)", " ", text)
    text = re.sub(r"[^a-z0-9&/+ ]+", " ", text)
    text = re.sub(r"\s+", " ", text)

    return text.strip()


# ---------------------------------------------------------------------------
# Input validation
# ---------------------------------------------------------------------------


def validate_inputs() -> None:
    """Check that all required stage inputs are available."""

    if not CHUNKS_FILE.is_file():
        raise FileNotFoundError(
            "Chunk file not found.\n"
            f"Expected: {CHUNKS_FILE}\n"
            "Run scripts/test_02_chunking.py first."
        )

    if not GROUND_TRUTH_FILE.is_file():
        raise FileNotFoundError(
            "Ground-truth workbook not found.\n"
            f"Expected: {GROUND_TRUTH_FILE}"
        )


def load_chunks() -> pd.DataFrame:
    """Load and validate the Stage 2 chunks."""

    chunks = pd.read_csv(CHUNKS_FILE, encoding="utf-8")
    required_columns = {"chunk_id", "text"}
    missing_columns = required_columns - set(chunks.columns)

    if missing_columns:
        raise ValueError(
            "The chunks file is missing required columns: "
            f"{sorted(missing_columns)}"
        )

    if chunks.empty:
        raise ValueError(f"The chunks file is empty: {CHUNKS_FILE}")

    chunks["text"] = chunks["text"].fillna("").astype(str)
    empty_chunks = int(chunks["text"].str.strip().eq("").sum())

    if empty_chunks:
        print(f"Warning: {empty_chunks} chunks contain empty text.")

    return chunks


# ---------------------------------------------------------------------------
# Ground-truth loading
# ---------------------------------------------------------------------------


def load_ground_truth() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Load strict terms, review terms, and negative benchmark terms."""

    strict_source = pd.read_excel(
        GROUND_TRUTH_FILE,
        sheet_name=GROUND_TRUTH_SHEET,
        engine="openpyxl",
    )

    negative_terms = pd.read_excel(
        GROUND_TRUTH_FILE,
        sheet_name=NEGATIVE_TERMS_SHEET,
        engine="openpyxl",
    )

    expected_gt_columns = {
        "gt_id",
        "term_canonical",
        "gt_category",
        "priority",
        "keep_for_tasks",
        "grounded_status_v2",
    }

    missing_gt_columns = expected_gt_columns - set(strict_source.columns)

    if missing_gt_columns:
        raise ValueError(
            f"{GROUND_TRUTH_SHEET} is missing columns: "
            f"{sorted(missing_gt_columns)}"
        )

    extraction_gt = strict_source[
        strict_source["keep_for_tasks"]
        .astype(str)
        .str.contains("term_extraction", case=False, na=False)
    ].copy()

    extraction_gt["normalized_gt"] = extraction_gt["term_canonical"].map(
        normalize_term
    )
    extraction_gt = extraction_gt[
        extraction_gt["normalized_gt"].ne("")
    ].copy()

    status = (
        extraction_gt["grounded_status_v2"]
        .fillna("")
        .astype(str)
        .str.casefold()
    )

    strict_gt = extraction_gt[status.ne(MANUAL_REVIEW_STATUS)].copy()
    review_gt = extraction_gt[status.eq(MANUAL_REVIEW_STATUS)].copy()

    if "candidate" not in negative_terms.columns:
        raise ValueError(
            f"{NEGATIVE_TERMS_SHEET} must contain a candidate column."
        )

    negative_terms["normalized_negative"] = negative_terms["candidate"].map(
        normalize_term
    )
    negative_terms = negative_terms[
        negative_terms["normalized_negative"].ne("")
    ].copy()

    return strict_gt, review_gt, negative_terms


# ---------------------------------------------------------------------------
# Candidate aggregation
# ---------------------------------------------------------------------------


def aggregate_candidates(
    raw_candidates: pd.DataFrame,
    method_name: str,
) -> pd.DataFrame:
    """Convert chunk-level occurrences into a unique ranked term inventory."""

    output_columns = [
        "method",
        "actual_extractor",
        "candidate",
        "normalized_candidate",
        "score",
        "global_rank",
        "best_local_rank",
        "occurrence_count",
        "chunk_count",
        "best_chunk_id",
        "source_sentence",
        "context_text",
    ]

    if raw_candidates.empty:
        return pd.DataFrame(columns=output_columns)

    work = raw_candidates.copy()
    required_columns = {
        "raw_term",
        "score",
        "rank",
        "extractor",
        "chunk_id",
        "source_sentence",
        "context_text",
    }
    missing_columns = required_columns - set(work.columns)

    if missing_columns:
        raise ValueError(
            f"{method_name} output is missing columns: "
            f"{sorted(missing_columns)}"
        )

    work["normalized_candidate"] = work["raw_term"].map(normalize_term)
    work = work[work["normalized_candidate"].ne("")].copy()
    work["score"] = pd.to_numeric(work["score"], errors="coerce").fillna(0.0)
    work["rank"] = pd.to_numeric(work["rank"], errors="coerce").fillna(999999)

    work = work.sort_values(
        by=["score", "rank", "normalized_candidate"],
        ascending=[False, True, True],
    )

    best_occurrences = (
        work.drop_duplicates(subset="normalized_candidate", keep="first")
        .set_index("normalized_candidate")
    )

    statistics = work.groupby("normalized_candidate").agg(
        occurrence_count=("normalized_candidate", "size"),
        chunk_count=("chunk_id", "nunique"),
        score=("score", "max"),
        best_local_rank=("rank", "min"),
    )

    unique_candidates = statistics.join(
        best_occurrences[
            [
                "raw_term",
                "extractor",
                "chunk_id",
                "source_sentence",
                "context_text",
            ]
        ]
    )

    unique_candidates = unique_candidates.reset_index().rename(
        columns={
            "raw_term": "candidate",
            "extractor": "actual_extractor",
            "chunk_id": "best_chunk_id",
        }
    )
    unique_candidates.insert(0, "method", method_name)

    unique_candidates = unique_candidates.sort_values(
        by=["score", "occurrence_count", "normalized_candidate"],
        ascending=[False, False, True],
    ).reset_index(drop=True)

    unique_candidates["global_rank"] = range(1, len(unique_candidates) + 1)
    return unique_candidates[output_columns]


# ---------------------------------------------------------------------------
# Evaluation
# ---------------------------------------------------------------------------


def _lookup_value(
    term: str,
    strict_lookup: pd.DataFrame,
    review_lookup: pd.DataFrame,
    column: str,
) -> object:
    """Safely retrieve a value from strict or review ground truth."""

    if term in strict_lookup.index and column in strict_lookup.columns:
        return strict_lookup.at[term, column]

    if term in review_lookup.index and column in review_lookup.columns:
        return review_lookup.at[term, column]

    return ""


def evaluate_candidates(
    candidates: pd.DataFrame,
    strict_gt: pd.DataFrame,
    review_gt: pd.DataFrame,
    negative_terms: pd.DataFrame,
    method_name: str,
) -> tuple[dict, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Evaluate unique candidates using normalized exact matching."""

    strict_lookup = (
        strict_gt.drop_duplicates("normalized_gt").set_index("normalized_gt")
    )
    review_lookup = (
        review_gt.drop_duplicates("normalized_gt").set_index("normalized_gt")
    )
    negative_lookup = (
        negative_terms.drop_duplicates("normalized_negative")
        .set_index("normalized_negative")
    )

    details = candidates.copy()

    if details.empty:
        details = details.copy()
        details["matched_gt_id"] = pd.Series(dtype="object")
        details["match_type"] = pd.Series(dtype="object")
        details["evaluation_flag"] = pd.Series(dtype="object")
        details["gt_category"] = pd.Series(dtype="object")
        details["expected_type"] = pd.Series(dtype="object")
        details["negative_benchmark_hit"] = pd.Series(dtype="bool")
        details["negative_reason"] = pd.Series(dtype="object")
    else:
        details["matched_gt_id"] = details["normalized_candidate"].map(
            lambda term: _lookup_value(
                term, strict_lookup, review_lookup, "gt_id"
            )
        )

        details["match_type"] = details["normalized_candidate"].map(
            lambda term: (
                "strict_exact"
                if term in strict_lookup.index
                else (
                    "manual_review_exact"
                    if term in review_lookup.index
                    else "unmatched"
                )
            )
        )

        details["evaluation_flag"] = details["match_type"].map(
            {
                "strict_exact": "TP",
                "manual_review_exact": "REVIEW",
                "unmatched": "FP",
            }
        )

        details["gt_category"] = details["normalized_candidate"].map(
            lambda term: _lookup_value(
                term, strict_lookup, review_lookup, "gt_category"
            )
        )

        details["expected_type"] = details["normalized_candidate"].map(
            lambda term: _lookup_value(
                term, strict_lookup, review_lookup, "expected_type"
            )
        )

        details["negative_benchmark_hit"] = details[
            "normalized_candidate"
        ].isin(negative_lookup.index)

        details["negative_reason"] = details["normalized_candidate"].map(
            lambda term: (
                negative_lookup.at[term, "reason_not_gt_domain_term"]
                if (
                    term in negative_lookup.index
                    and "reason_not_gt_domain_term" in negative_lookup.columns
                )
                else ""
            )
        )

    strict_prediction_rows = details[
        details["match_type"].ne("manual_review_exact")
    ]
    predicted_terms = set(strict_prediction_rows["normalized_candidate"])
    gold_terms = set(strict_gt["normalized_gt"])

    true_positive_terms = predicted_terms & gold_terms
    false_positive_terms = predicted_terms - gold_terms
    false_negative_terms = gold_terms - predicted_terms

    true_positives = len(true_positive_terms)
    false_positives = len(false_positive_terms)
    false_negatives = len(false_negative_terms)

    precision_denominator = true_positives + false_positives
    recall_denominator = true_positives + false_negatives

    precision = (
        true_positives / precision_denominator
        if precision_denominator
        else 0.0
    )
    recall = (
        true_positives / recall_denominator
        if recall_denominator
        else 0.0
    )
    f1 = (
        2 * precision * recall / (precision + recall)
        if precision + recall
        else 0.0
    )

    metric = {
        "method": method_name,
        "candidate_count": len(details),
        "strict_scored_prediction_count": len(strict_prediction_rows),
        "strict_gold_count": len(gold_terms),
        "manual_review_gold_count": review_gt["normalized_gt"].nunique(),
        "true_positives": true_positives,
        "false_positives": false_positives,
        "false_negatives": false_negatives,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "manual_review_matches": int(
            details["match_type"].eq("manual_review_exact").sum()
        ),
        "negative_term_hits": int(details["negative_benchmark_hit"].sum()),
    }

    gt_weights = (
        strict_gt["priority"]
        .fillna("low")
        .astype(str)
        .str.casefold()
        .map(PRIORITY_WEIGHTS)
        .fillna(1.0)
    )
    matched_gt_mask = strict_gt["normalized_gt"].isin(predicted_terms)
    total_weight = float(gt_weights.sum())
    matched_weight = float(gt_weights[matched_gt_mask].sum())
    metric["priority_weighted_recall"] = (
        matched_weight / total_weight if total_weight else 0.0
    )

    false_negative_report = strict_gt[
        strict_gt["normalized_gt"].isin(false_negative_terms)
    ].copy()

    selected_fn_columns = [
        column
        for column in [
            "gt_id",
            "term_canonical",
            "normalized_gt",
            "gt_category",
            "expected_type",
            "priority",
            "grounded_status_v2",
            "evidence_chapter",
            "evidence_page_pdf",
            "evidence_snippet",
        ]
        if column in false_negative_report.columns
    ]

    false_negative_report = false_negative_report[selected_fn_columns].copy()
    false_negative_report.insert(0, "method", method_name)
    false_negative_report["evaluation_flag"] = "FN"

    category_rows = []
    for category, category_group in strict_gt.groupby("gt_category"):
        category_gold = set(category_group["normalized_gt"])
        category_matches = category_gold & predicted_terms
        category_rows.append(
            {
                "method": method_name,
                "gt_category": category,
                "gold_count": len(category_gold),
                "matched_count": len(category_matches),
                "missed_count": len(category_gold) - len(category_matches),
                "recall": (
                    len(category_matches) / len(category_gold)
                    if category_gold
                    else 0.0
                ),
            }
        )

    category_report = pd.DataFrame(category_rows)
    negative_hits = details[details["negative_benchmark_hit"]].copy()

    return (
        metric,
        details,
        false_negative_report,
        category_report,
        negative_hits,
    )


# ---------------------------------------------------------------------------
# Output helpers
# ---------------------------------------------------------------------------


def save_method_results(
    method_directory: Path,
    raw_candidates: pd.DataFrame,
    unique_candidates: pd.DataFrame,
    evaluation_details: pd.DataFrame,
    false_negatives: pd.DataFrame,
    category_report: pd.DataFrame,
    negative_hits: pd.DataFrame,
    metric: dict,
) -> None:
    """Save every artifact for one extraction method."""

    method_directory.mkdir(parents=True, exist_ok=True)

    raw_candidates.to_csv(
        method_directory / "candidates_all_occurrences.csv",
        index=False,
        encoding="utf-8",
    )
    unique_candidates.to_csv(
        method_directory / "candidates_unique_ranked.csv",
        index=False,
        encoding="utf-8",
    )
    evaluation_details.to_csv(
        method_directory / "evaluation_predictions.csv",
        index=False,
        encoding="utf-8",
    )
    false_negatives.to_csv(
        method_directory / "false_negatives.csv",
        index=False,
        encoding="utf-8",
    )
    category_report.to_csv(
        method_directory / "category_recall.csv",
        index=False,
        encoding="utf-8",
    )
    negative_hits.to_csv(
        method_directory / "negative_term_hits.csv",
        index=False,
        encoding="utf-8",
    )
    pd.DataFrame([metric]).to_csv(
        method_directory / "metrics.csv",
        index=False,
        encoding="utf-8",
    )


def format_excel_report(workbook_path: Path) -> None:
    """Apply basic formatting to the combined Excel report."""

    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("Warning: openpyxl formatting was skipped.")
        return

    workbook = load_workbook(workbook_path)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False

        if worksheet.max_row > 0 and worksheet.max_column > 0:
            worksheet.auto_filter.ref = worksheet.dimensions

        # Corrected block: iterate over every cell in the first row.
        for header_cell in worksheet[1]:
            header_cell.fill = header_fill
            header_cell.font = header_font
            header_cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

        worksheet.row_dimensions[1].height = 30

        for column_number in range(1, worksheet.max_column + 1):
            column_letter = get_column_letter(column_number)
            values = []

            for row_number in range(1, min(worksheet.max_row, 200) + 1):
                cell_value = worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).value
                if cell_value is not None:
                    values.append(str(cell_value))

            maximum_length = max(
                (len(value) for value in values),
                default=10,
            )
            worksheet.column_dimensions[column_letter].width = min(
                max(maximum_length + 2, 12),
                45,
            )

        if worksheet.title == "Metrics Summary":
            percentage_headers = {
                "precision",
                "recall",
                "f1",
                "priority_weighted_recall",
            }

            header_positions = {
                worksheet.cell(row=1, column=column).value: column
                for column in range(1, worksheet.max_column + 1)
            }

            for header in percentage_headers:
                if header not in header_positions:
                    continue

                column_number = header_positions[header]
                for row_number in range(2, worksheet.max_row + 1):
                    worksheet.cell(
                        row=row_number,
                        column=column_number,
                    ).number_format = "0.00%"

    workbook.save(workbook_path)


# ---------------------------------------------------------------------------
# Main extraction experiment
# ---------------------------------------------------------------------------


def main() -> None:
    validate_inputs()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    methods_directory = OUTPUT_DIR / "by_method"
    methods_directory.mkdir(parents=True, exist_ok=True)

    print("Loading chunks")
    chunks = load_chunks()

    print("Loading ground truth")
    strict_gt, review_gt, negative_terms = load_ground_truth()

    print()
    print(f"Chunks: {len(chunks)}")
    print(f"Strict ground-truth terms: {strict_gt['normalized_gt'].nunique()}")
    print(
        "Manual-review ground-truth terms: "
        f"{review_gt['normalized_gt'].nunique()}"
    )
    print(
        "Negative benchmark terms: "
        f"{negative_terms['normalized_negative'].nunique()}"
    )

    extraction_methods = [
        {
            "name": "yake",
            "function": yake_extract,
            "parameters": {"top_k": None},
        },
        {
            "name": "tfidf",
            "function": tfidf,
            "parameters": {"top_k": None},
        },
        {
            "name": "noun_phrases",
            "function": noun_phrases,
            "parameters": {"top_k": None},
        },
        {
            "name": "textrank",
            "function": textrank,
            "parameters": {"top_k": None},
        },
        {
            "name": "keybert_minilm",
            "function": keybert_extract,
            "parameters": {
                "top_k": None,
                "model": "all-MiniLM-L6-v2",
            },
        },
        {
            "name": "keybert_mpnet",
            "function": keybert_extract,
            "parameters": {
                "top_k": None,
                "model": "all-mpnet-base-v2",
            },
        },
    ]

    metrics = []
    combined_candidates = []
    combined_evaluations = []
    combined_false_negatives = []
    combined_category_reports = []
    combined_negative_hits = []

    manifest = {
        "stage": "03_non_llm_extraction",
        "chunks_file": str(CHUNKS_FILE),
        "ground_truth_file": str(GROUND_TRUTH_FILE),
        "output_directory": str(OUTPUT_DIR),
        "candidate_selection_rule": CANDIDATE_SELECTION_RULE,
        "chunk_count": len(chunks),
        "strict_ground_truth_count": int(
            strict_gt["normalized_gt"].nunique()
        ),
        "manual_review_ground_truth_count": int(
            review_gt["normalized_gt"].nunique()
        ),
        "matching_rule": "case-insensitive normalized exact match",
        "ranked_cutoff_metrics": False,
        "methods": [],
    }

    for method_configuration in extraction_methods:
        method_name = method_configuration["name"]
        extraction_function = method_configuration["function"]
        parameters = method_configuration["parameters"]
        method_directory = methods_directory / method_name
        method_directory.mkdir(parents=True, exist_ok=True)
        start_time = time.time()

        print()
        print("=" * 80)
        print(f"Running: {method_name}")
        print(f"Parameters: {parameters}")
        print("=" * 80)

        try:
            raw_candidates = extraction_function(chunks, **parameters)
            unique_candidates = aggregate_candidates(
                raw_candidates,
                method_name,
            )

            (
                metric,
                evaluation_details,
                false_negatives,
                category_report,
                negative_hits,
            ) = evaluate_candidates(
                candidates=unique_candidates,
                strict_gt=strict_gt,
                review_gt=review_gt,
                negative_terms=negative_terms,
                method_name=method_name,
            )

            save_method_results(
                method_directory=method_directory,
                raw_candidates=raw_candidates,
                unique_candidates=unique_candidates,
                evaluation_details=evaluation_details,
                false_negatives=false_negatives,
                category_report=category_report,
                negative_hits=negative_hits,
                metric=metric,
            )

            metrics.append(metric)
            combined_candidates.append(unique_candidates)
            combined_evaluations.append(evaluation_details)
            combined_false_negatives.append(false_negatives)
            combined_category_reports.append(category_report)

            if not negative_hits.empty:
                combined_negative_hits.append(negative_hits)

            elapsed_time = round(time.time() - start_time, 3)
            actual_extractors = sorted(
                raw_candidates["extractor"]
                .dropna()
                .astype(str)
                .unique()
                .tolist()
            )

            manifest["methods"].append(
                {
                    "method": method_name,
                    "status": "completed",
                    "parameters": parameters,
                    "actual_extractors": actual_extractors,
                    "duration_seconds": elapsed_time,
                    "raw_candidate_occurrences": len(raw_candidates),
                    "unique_candidates": len(unique_candidates),
                    "precision": metric["precision"],
                    "recall": metric["recall"],
                    "f1": metric["f1"],
                }
            )

            print(f"Completed {method_name}")
            print(f"Raw occurrences: {len(raw_candidates)}")
            print(f"Unique candidates: {len(unique_candidates)}")
            print(f"Precision: {metric['precision']:.4f}")
            print(f"Recall: {metric['recall']:.4f}")
            print(f"F1: {metric['f1']:.4f}")
            print(f"Duration: {elapsed_time} seconds")

            if any(
                "fallback" in extractor.casefold()
                for extractor in actual_extractors
            ):
                print(
                    "Warning: this method used a fallback extractor. "
                    "Check installed dependencies."
                )

        except Exception as error:
            elapsed_time = round(time.time() - start_time, 3)
            error_details = traceback.format_exc()
            error_file = method_directory / "ERROR.txt"
            error_file.write_text(error_details, encoding="utf-8")

            manifest["methods"].append(
                {
                    "method": method_name,
                    "status": "failed",
                    "parameters": parameters,
                    "duration_seconds": elapsed_time,
                    "error": str(error),
                    "error_file": str(error_file),
                }
            )

            print(f"Failed: {method_name}")
            print(f"Reason: {error}")
            print(f"Details saved to: {error_file}")

    if metrics:
        metrics_summary = (
            pd.DataFrame(metrics)
            .sort_values(
                by=["f1", "recall", "precision"],
                ascending=[False, False, False],
            )
            .reset_index(drop=True)
        )
    else:
        metrics_summary = pd.DataFrame()

    all_candidates = (
        pd.concat(combined_candidates, ignore_index=True)
        if combined_candidates
        else pd.DataFrame()
    )
    all_evaluations = (
        pd.concat(combined_evaluations, ignore_index=True)
        if combined_evaluations
        else pd.DataFrame()
    )
    all_false_negatives = (
        pd.concat(combined_false_negatives, ignore_index=True)
        if combined_false_negatives
        else pd.DataFrame()
    )
    all_category_reports = (
        pd.concat(combined_category_reports, ignore_index=True)
        if combined_category_reports
        else pd.DataFrame()
    )
    all_negative_hits = (
        pd.concat(combined_negative_hits, ignore_index=True)
        if combined_negative_hits
        else pd.DataFrame()
    )

    metrics_summary.to_csv(
        OUTPUT_DIR / "metrics_summary.csv",
        index=False,
        encoding="utf-8",
    )
    all_candidates.to_csv(
        OUTPUT_DIR / "all_methods_candidates.csv",
        index=False,
        encoding="utf-8",
    )
    all_evaluations.to_csv(
        OUTPUT_DIR / "all_methods_evaluation.csv",
        index=False,
        encoding="utf-8",
    )
    all_false_negatives.to_csv(
        OUTPUT_DIR / "all_methods_false_negatives.csv",
        index=False,
        encoding="utf-8",
    )
    all_category_reports.to_csv(
        OUTPUT_DIR / "all_methods_category_recall.csv",
        index=False,
        encoding="utf-8",
    )
    all_negative_hits.to_csv(
        OUTPUT_DIR / "all_methods_negative_hits.csv",
        index=False,
        encoding="utf-8",
    )

    report_file = OUTPUT_DIR / "non_llm_extraction_report.xlsx"
    readme = pd.DataFrame(
        {
            "Field": [
                "Stage",
                "Purpose",
                "Input chunks",
                "Ground-truth sheet",
                "Strict GT count",
                "Manual-review GT count",
                "Matching rule",
                "Candidate-selection rule",
                "Important interpretation",
            ],
            "Value": [
                "03 Non-LLM extraction",
                "Compare non-LLM term extraction methods against S1 ground truth.",
                str(CHUNKS_FILE),
                GROUND_TRUTH_SHEET,
                strict_gt["normalized_gt"].nunique(),
                review_gt["normalized_gt"].nunique(),
                "Case-insensitive normalized exact match.",
                CANDIDATE_SELECTION_RULE,
                (
                    "An unmatched candidate is not automatically an invalid "
                    "domain term. The ground truth may not enumerate every "
                    "valid technical term."
                ),
            ],
        }
    )

    with pd.ExcelWriter(report_file, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        metrics_summary.to_excel(
            writer,
            sheet_name="Metrics Summary",
            index=False,
        )
        all_category_reports.to_excel(
            writer,
            sheet_name="Category Recall",
            index=False,
        )
        all_evaluations.to_excel(
            writer,
            sheet_name="Prediction Evaluation",
            index=False,
        )
        all_false_negatives.to_excel(
            writer,
            sheet_name="False Negatives",
            index=False,
        )
        all_negative_hits.to_excel(
            writer,
            sheet_name="Negative Hits",
            index=False,
        )
        all_candidates.to_excel(
            writer,
            sheet_name="All Candidates",
            index=False,
        )

    format_excel_report(report_file)

    manifest_file = OUTPUT_DIR / "run_manifest.json"
    manifest_file.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print()
    print("=" * 80)
    print("Stage 03 extraction completed")
    print("=" * 80)
    print()
    print(f"Metrics summary: {OUTPUT_DIR / 'metrics_summary.csv'}")
    print(f"Excel report: {report_file}")
    print(f"Manifest: {manifest_file}")
    print(f"Method outputs: {methods_directory}")

    if not metrics_summary.empty:
        print()
        print("Method comparison:")
        print(
            metrics_summary[
                [
                    "method",
                    "candidate_count",
                    "true_positives",
                    "precision",
                    "recall",
                    "f1",
                    "negative_term_hits",
                ]
            ].to_string(index=False)
        )


if __name__ == "__main__":
    main()
