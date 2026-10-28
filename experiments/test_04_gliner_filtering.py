from __future__ import annotations

import json
import re
import time
import traceback
from pathlib import Path

import pandas as pd
from gliner import GLiNER


PROJECT_DIR = Path("ontology_experiment_framework")
INPUT_FILE = PROJECT_DIR / "stage_outputs/03c_weighted_voting/best_f1_candidates.csv"
GROUND_TRUTH_FILE = PROJECT_DIR / "data/Ground_Truth.xlsx"
OUTPUT_DIR = PROJECT_DIR / "stage_outputs/04_gliner_filtering"

MODEL_NAME = "urchade/gliner_medium-v2.1"
INFERENCE_THRESHOLD = 0.20
TEST_THRESHOLDS = [0.20, 0.25, 0.30, 0.35, 0.40, 0.45, 0.50, 0.55, 0.60, 0.65, 0.70]

GROUND_TRUTH_SHEET = "S1_GT_TERMS"
NEGATIVE_TERMS_SHEET = "S7_NEGATIVE_TERMS"
MANUAL_REVIEW_STATUS = "needs_manual_review"
PRIORITY_WEIGHTS = {"high": 3.0, "medium": 2.0, "low": 1.0}

# GLiNER is zero-shot: labels are supplied at inference time.
ALL_NEGATIVE_LABELS = [
    "person",
    "author",
    "organization",
    "company",
    "government agency",
    "country",
    "city",
    "state",
    "location",
    "university",
    "publication",
    "book",
    "journal",
    "conference",
    "date",
]

# Conservative policy avoids automatically dropping every organization/location.
CONSERVATIVE_LABELS = {
    "person",
    "author",
    "country",
    "city",
    "state",
    "university",
    "publication",
    "book",
    "journal",
    "conference",
    "date",
}

# Strict policy also drops organizations, agencies, companies, and locations.
STRICT_LABELS = set(ALL_NEGATIVE_LABELS)

POLICIES = {
    "conservative": CONSERVATIVE_LABELS,
    "strict": STRICT_LABELS,
}


def normalize_term(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    text = str(value).casefold().strip()
    for dash in ("‐", "‑", "‒", "–", "—"):
        text = text.replace(dash, "-")
    text = re.sub(r"(?<=\w)-(?=\w)", " ", text)
    text = re.sub(r"[^a-z0-9&/+ ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def safe_text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def validate_inputs() -> None:
    missing = []
    if not INPUT_FILE.is_file():
        missing.append(str(INPUT_FILE))
    if not GROUND_TRUTH_FILE.is_file():
        missing.append(str(GROUND_TRUTH_FILE))
    if missing:
        raise FileNotFoundError("Missing required files:\n" + "\n".join(missing))


def load_candidates() -> pd.DataFrame:
    candidates = pd.read_csv(INPUT_FILE, encoding="utf-8")
    required = {"candidate", "normalized_candidate"}
    missing = required - set(candidates.columns)
    if missing:
        raise ValueError(f"{INPUT_FILE} is missing columns: {sorted(missing)}")

    for column in ("source_sentence", "context_text"):
        if column not in candidates.columns:
            candidates[column] = ""

    candidates = candidates.copy()
    candidates["candidate"] = candidates["candidate"].map(safe_text)
    candidates["normalized_candidate"] = candidates["normalized_candidate"].map(normalize_term)
    candidates["source_sentence"] = candidates["source_sentence"].map(safe_text)
    candidates["context_text"] = candidates["context_text"].map(safe_text)
    candidates = candidates[candidates["normalized_candidate"].ne("")]
    candidates = candidates.drop_duplicates("normalized_candidate", keep="first").reset_index(drop=True)
    if candidates.empty:
        raise ValueError("The Stage 03c candidate inventory is empty.")
    return candidates


def load_ground_truth():
    gt = pd.read_excel(GROUND_TRUTH_FILE, sheet_name=GROUND_TRUTH_SHEET, engine="openpyxl")
    negatives = pd.read_excel(GROUND_TRUTH_FILE, sheet_name=NEGATIVE_TERMS_SHEET, engine="openpyxl")

    gt = gt[
        gt["keep_for_tasks"].astype(str).str.contains("term_extraction", case=False, na=False)
    ].copy()
    gt["normalized_gt"] = gt["term_canonical"].map(normalize_term)
    gt = gt[gt["normalized_gt"].ne("")].copy()
    status = gt["grounded_status_v2"].fillna("").astype(str).str.casefold()
    strict_gt = gt[status.ne(MANUAL_REVIEW_STATUS)].copy()
    review_gt = gt[status.eq(MANUAL_REVIEW_STATUS)].copy()

    negatives["normalized_negative"] = negatives["candidate"].map(normalize_term)
    negatives = negatives[negatives["normalized_negative"].ne("")].copy()
    return strict_gt, review_gt, negatives


def make_gliner_text(candidate: str, source_sentence: str) -> tuple[str, int, int]:
    prefix = "Candidate: "
    text = prefix + candidate
    candidate_start = len(prefix)
    candidate_end = candidate_start + len(candidate)
    if source_sentence:
        text += ". Source sentence: " + source_sentence
    return text, candidate_start, candidate_end


def spans_overlap(start_a: int, end_a: int, start_b: int, end_b: int) -> bool:
    return max(start_a, start_b) < min(end_a, end_b)


def run_gliner(candidates: pd.DataFrame, model: GLiNER) -> pd.DataFrame:
    rows = []
    total = len(candidates)

    for index, row in enumerate(candidates.itertuples(), start=1):
        text, candidate_start, candidate_end = make_gliner_text(
            safe_text(row.candidate),
            safe_text(row.source_sentence),
        )

        entities = model.predict_entities(
            text,
            ALL_NEGATIVE_LABELS,
            threshold=INFERENCE_THRESHOLD,
        )

        overlapping = [
            entity
            for entity in entities
            if spans_overlap(
                int(entity["start"]),
                int(entity["end"]),
                candidate_start,
                candidate_end,
            )
        ]
        overlapping.sort(key=lambda entity: float(entity.get("score", 0.0)), reverse=True)

        best_entity = overlapping[0] if overlapping else None
        output = row._asdict()
        output["gliner_input_text"] = text
        output["gliner_entity_count_on_candidate"] = len(overlapping)
        output["gliner_best_label"] = best_entity.get("label", "") if best_entity else ""
        output["gliner_best_score"] = float(best_entity.get("score", 0.0)) if best_entity else 0.0
        output["gliner_best_span_text"] = best_entity.get("text", "") if best_entity else ""
        output["gliner_all_candidate_entities"] = json.dumps(overlapping, ensure_ascii=False)
        rows.append(output)

        if index % 100 == 0 or index == total:
            print(f"GLiNER processed {index}/{total}")

    result = pd.DataFrame(rows)
    # itertuples adds Index only if index=True; remove defensively.
    result = result.drop(columns=["Index"], errors="ignore")
    return result


def annotate_matches(candidates, strict_gt, review_gt, negatives):
    strict_lookup = strict_gt.drop_duplicates("normalized_gt").set_index("normalized_gt")
    review_lookup = review_gt.drop_duplicates("normalized_gt").set_index("normalized_gt")
    negative_set = set(negatives["normalized_negative"].dropna())

    result = candidates.copy()
    result["match_type"] = result["normalized_candidate"].map(
        lambda term: "strict_exact" if term in strict_lookup.index else (
            "manual_review_exact" if term in review_lookup.index else "unmatched"
        )
    )
    result["evaluation_flag_before_filter"] = result["match_type"].map(
        {"strict_exact": "TP", "manual_review_exact": "REVIEW", "unmatched": "FP"}
    )
    result["matched_gt_id"] = result["normalized_candidate"].map(
        lambda term: strict_lookup.at[term, "gt_id"] if term in strict_lookup.index else (
            review_lookup.at[term, "gt_id"] if term in review_lookup.index else ""
        )
    )
    result["negative_benchmark_hit"] = result["normalized_candidate"].isin(negative_set)
    return result


def evaluate_filter(experiment, annotated, drop_mask, strict_gt):
    retained = annotated[~drop_mask].copy()
    dropped = annotated[drop_mask].copy()

    strict_unique = strict_gt.drop_duplicates("normalized_gt").copy()
    gold = set(strict_unique["normalized_gt"])
    scored = retained[retained["match_type"].ne("manual_review_exact")]
    predicted = set(scored["normalized_candidate"])

    tp = len(predicted & gold)
    fp = len(predicted - gold)
    fn = len(gold - predicted)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0

    gt_weights = (
        strict_unique["priority"].fillna("low").astype(str).str.casefold()
        .map(PRIORITY_WEIGHTS).fillna(1.0)
    )
    matched_mask = strict_unique["normalized_gt"].isin(predicted)
    weighted_recall = float(gt_weights[matched_mask].sum() / gt_weights.sum()) if gt_weights.sum() else 0.0

    metric = {
        "experiment": experiment,
        "input_candidate_count": len(annotated),
        "retained_candidate_count": len(retained),
        "dropped_candidate_count": len(dropped),
        "true_positives": tp,
        "false_positives": fp,
        "false_negatives": fn,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "priority_weighted_recall": weighted_recall,
        "tp_retained": int(retained["match_type"].eq("strict_exact").sum()),
        "tp_dropped": int(dropped["match_type"].eq("strict_exact").sum()),
        "fp_retained": int(retained["match_type"].eq("unmatched").sum()),
        "fp_dropped": int(dropped["match_type"].eq("unmatched").sum()),
        "manual_review_retained": int(retained["match_type"].eq("manual_review_exact").sum()),
        "negative_hits_retained": int(retained["negative_benchmark_hit"].sum()),
        "negative_hits_dropped": int(dropped["negative_benchmark_hit"].sum()),
    }

    retained["gliner_experiment"] = experiment
    retained["gliner_decision"] = "retained"
    dropped["gliner_experiment"] = experiment
    dropped["gliner_decision"] = "dropped"
    return metric, retained, dropped


def format_workbook(path: Path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        if worksheet.max_row > 0 and worksheet.max_column > 0:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for column in range(1, worksheet.max_column + 1):
            values = [
                len(str(worksheet.cell(row, column).value or ""))
                for row in range(1, min(worksheet.max_row, 150) + 1)
            ]
            worksheet.column_dimensions[get_column_letter(column)].width = min(
                max(max(values, default=10) + 2, 11), 42
            )
        if worksheet.title == "Metrics Summary":
            headers = {worksheet.cell(1, column).value: column for column in range(1, worksheet.max_column + 1)}
            for name in ("precision", "recall", "f1", "priority_weighted_recall"):
                if name in headers:
                    for row in range(2, worksheet.max_row + 1):
                        worksheet.cell(row, headers[name]).number_format = "0.00%"
    workbook.save(path)


def main():
    validate_inputs()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    by_filter = OUTPUT_DIR / "by_filter"
    by_filter.mkdir(parents=True, exist_ok=True)

    candidates = load_candidates()
    strict_gt, review_gt, negatives = load_ground_truth()

    print(f"Loading GLiNER model: {MODEL_NAME}")
    started = time.time()
    model = GLiNER.from_pretrained(MODEL_NAME)
    predictions = run_gliner(candidates, model)
    annotated = annotate_matches(predictions, strict_gt, review_gt, negatives)
    annotated.to_csv(OUTPUT_DIR / "gliner_scored_candidates.csv", index=False, encoding="utf-8")

    metrics = []
    outputs = {}

    baseline_metric, baseline_retained, baseline_dropped = evaluate_filter(
        "baseline_no_gliner_filter",
        annotated,
        pd.Series(False, index=annotated.index),
        strict_gt,
    )
    metrics.append(baseline_metric)
    outputs[baseline_metric["experiment"]] = (baseline_retained, baseline_dropped)

    for policy_name, policy_labels in POLICIES.items():
        label_match = annotated["gliner_best_label"].isin(policy_labels)

        for threshold in TEST_THRESHOLDS:
            experiment = f"{policy_name}__score_ge_{threshold:.2f}"
            drop_mask = label_match & (annotated["gliner_best_score"] >= threshold)
            metric, retained, dropped = evaluate_filter(
                experiment,
                annotated,
                drop_mask,
                strict_gt,
            )
            metrics.append(metric)
            outputs[experiment] = (retained, dropped)

    metrics_df = pd.DataFrame(metrics).sort_values(
        ["f1", "precision", "recall"], ascending=[False, False, False]
    ).reset_index(drop=True)
    metrics_df.to_csv(OUTPUT_DIR / "gliner_metrics.csv", index=False, encoding="utf-8")

    best_experiment = metrics_df.iloc[0]["experiment"]
    best_retained, best_dropped = outputs[best_experiment]
    best_retained.to_csv(OUTPUT_DIR / "best_gliner_retained.csv", index=False, encoding="utf-8")
    best_dropped.to_csv(OUTPUT_DIR / "best_gliner_dropped.csv", index=False, encoding="utf-8")

    for experiment, (retained, dropped) in outputs.items():
        safe_name = experiment.replace(".", "_")
        retained.to_csv(by_filter / f"{safe_name}__retained.csv", index=False, encoding="utf-8")
        dropped.to_csv(by_filter / f"{safe_name}__dropped.csv", index=False, encoding="utf-8")

    label_summary = (
        annotated.groupby("gliner_best_label", dropna=False)
        .agg(
            candidate_count=("normalized_candidate", "size"),
            average_score=("gliner_best_score", "mean"),
            maximum_score=("gliner_best_score", "max"),
            strict_gt_count=("match_type", lambda values: int((values == "strict_exact").sum())),
            unmatched_count=("match_type", lambda values: int((values == "unmatched").sum())),
        )
        .reset_index()
        .sort_values("candidate_count", ascending=False)
    )
    label_summary.to_csv(OUTPUT_DIR / "gliner_label_summary.csv", index=False, encoding="utf-8")

    report = OUTPUT_DIR / "gliner_filtering_report.xlsx"
    readme = pd.DataFrame({
        "Field": [
            "Purpose", "Input", "Model", "Inference threshold", "Policies",
            "Best experiment", "Important caution",
        ],
        "Value": [
            "Remove unwanted named-entity candidates before seed-similarity filtering.",
            str(INPUT_FILE),
            MODEL_NAME,
            INFERENCE_THRESHOLD,
            "Conservative and strict label sets across thresholds 0.20 to 0.70.",
            best_experiment,
            "Threshold and policy selection on the same GT is a development result.",
        ],
    })
    with pd.ExcelWriter(report, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        metrics_df.to_excel(writer, sheet_name="Metrics Summary", index=False)
        label_summary.to_excel(writer, sheet_name="Label Summary", index=False)
        best_retained.to_excel(writer, sheet_name="Best Retained", index=False)
        best_dropped.to_excel(writer, sheet_name="Best Dropped", index=False)
    format_workbook(report)

    manifest = {
        "stage": "04_gliner_filtering",
        "input_file": str(INPUT_FILE),
        "model": MODEL_NAME,
        "labels": ALL_NEGATIVE_LABELS,
        "policies": {name: sorted(labels) for name, labels in POLICIES.items()},
        "inference_threshold": INFERENCE_THRESHOLD,
        "tested_thresholds": TEST_THRESHOLDS,
        "best_experiment_by_f1": best_experiment,
        "best_f1": float(metrics_df.iloc[0]["f1"]),
        "duration_seconds": round(time.time() - started, 3),
        "next_stage_input": str(OUTPUT_DIR / "best_gliner_retained.csv"),
    }
    (OUTPUT_DIR / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print("\n" + "=" * 80)
    print("Stage 04 GLiNER filtering completed")
    print("=" * 80)
    print(metrics_df.head(15).to_string(index=False))
    print(f"\nBest experiment: {best_experiment}")
    print(f"Best retained file: {OUTPUT_DIR / 'best_gliner_retained.csv'}")
    print(f"Excel report: {report}")


if __name__ == "__main__":
    main()
