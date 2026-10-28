from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd


PROJECT_DIR = Path("ontology_experiment_framework")
EXTRACTION_DIR = PROJECT_DIR / "stage_outputs/03_non_llm_extraction"
GROUND_TRUTH_FILE = PROJECT_DIR / "data/Ground_Truth.xlsx"
OUTPUT_DIR = PROJECT_DIR / "stage_outputs/03b_extraction_ensembles"

GROUND_TRUTH_SHEET = "S1_GT_TERMS"
NEGATIVE_TERMS_SHEET = "S7_NEGATIVE_TERMS"
MANUAL_REVIEW_STATUS = "needs_manual_review"

PRIORITY_WEIGHTS = {
    "high": 3.0,
    "medium": 2.0,
    "low": 1.0,
}

# These are the three complementary methods used by the ensemble experiment.
METHOD_FILES = {
    "yake": EXTRACTION_DIR / "by_method/yake/candidates_unique_ranked.csv",
    "tfidf": EXTRACTION_DIR / "by_method/tfidf/candidates_unique_ranked.csv",
    "mpnet": EXTRACTION_DIR / "by_method/keybert_mpnet/candidates_unique_ranked.csv",
}


def normalize_term(value: object) -> str:
    if value is None or pd.isna(value):
        return ""

    text = str(value).casefold().strip()
    text = (
        text.replace("‐", "-")
        .replace("‑", "-")
        .replace("‒", "-")
        .replace("–", "-")
        .replace("—", "-")
    )
    text = re.sub(r"(?<=\w)-(?=\w)", " ", text)
    text = re.sub(r"[^a-z0-9&/+ ]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def validate_inputs() -> None:
    if not GROUND_TRUTH_FILE.is_file():
        raise FileNotFoundError(f"Ground truth not found: {GROUND_TRUTH_FILE}")

    missing = [str(path) for path in METHOD_FILES.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError(
            "Required Stage 03 outputs are missing:\n" + "\n".join(missing)
        )


def load_ground_truth() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    source = pd.read_excel(
        GROUND_TRUTH_FILE,
        sheet_name=GROUND_TRUTH_SHEET,
        engine="openpyxl",
    )
    negatives = pd.read_excel(
        GROUND_TRUTH_FILE,
        sheet_name=NEGATIVE_TERMS_SHEET,
        engine="openpyxl",
    )

    source = source[
        source["keep_for_tasks"]
        .astype(str)
        .str.contains("term_extraction", case=False, na=False)
    ].copy()
    source["normalized_gt"] = source["term_canonical"].map(normalize_term)
    source = source[source["normalized_gt"].ne("")].copy()

    status = source["grounded_status_v2"].fillna("").astype(str).str.casefold()
    strict_gt = source[status.ne(MANUAL_REVIEW_STATUS)].copy()
    review_gt = source[status.eq(MANUAL_REVIEW_STATUS)].copy()

    negatives["normalized_negative"] = negatives["candidate"].map(normalize_term)
    negatives = negatives[negatives["normalized_negative"].ne("")].copy()
    return strict_gt, review_gt, negatives


def load_method_candidates() -> tuple[dict[str, pd.DataFrame], pd.DataFrame]:
    frames = {}
    combined = []

    for method, path in METHOD_FILES.items():
        frame = pd.read_csv(path, encoding="utf-8")

        required = {"candidate", "normalized_candidate", "score", "global_rank"}
        missing = required - set(frame.columns)
        if missing:
            raise ValueError(f"{path} is missing columns: {sorted(missing)}")

        frame = frame.copy()
        frame["normalized_candidate"] = frame["normalized_candidate"].map(
            normalize_term
        )
        frame = frame[frame["normalized_candidate"].ne("")].copy()
        frame = frame.sort_values("global_rank").drop_duplicates(
            "normalized_candidate", keep="first"
        )
        frames[method] = frame

        selected = frame[
            ["normalized_candidate", "candidate", "score", "global_rank"]
        ].copy()
        selected["source_method"] = method
        combined.append(selected)

    return frames, pd.concat(combined, ignore_index=True)


def create_support_table(
    method_frames: dict[str, pd.DataFrame],
    combined: pd.DataFrame,
) -> pd.DataFrame:
    all_terms = sorted(
        set().union(
            *(set(frame["normalized_candidate"]) for frame in method_frames.values())
        )
    )
    support = pd.DataFrame({"normalized_candidate": all_terms})

    for method, frame in method_frames.items():
        lookup = frame.set_index("normalized_candidate")
        support[f"found_by_{method}"] = support["normalized_candidate"].isin(
            lookup.index
        )
        support[f"{method}_rank"] = support["normalized_candidate"].map(
            lookup["global_rank"]
        )
        support[f"{method}_score"] = support["normalized_candidate"].map(
            lookup["score"]
        )

    support["vote_count"] = support[
        ["found_by_yake", "found_by_tfidf", "found_by_mpnet"]
    ].sum(axis=1)

    # Prefer the surface form from YAKE, then MPNet, then TF-IDF.
    surface_maps = {
        method: frame.set_index("normalized_candidate")["candidate"].to_dict()
        for method, frame in method_frames.items()
    }
    support["candidate"] = support["normalized_candidate"].map(
        lambda term: surface_maps["yake"].get(
            term,
            surface_maps["mpnet"].get(term, surface_maps["tfidf"].get(term, term)),
        )
    )

    rank_columns = ["yake_rank", "tfidf_rank", "mpnet_rank"]
    support["best_source_rank"] = support[rank_columns].min(axis=1, skipna=True)
    support["average_source_rank"] = support[rank_columns].mean(axis=1, skipna=True)
    return support


def build_ensembles(support: pd.DataFrame) -> dict[str, pd.DataFrame]:
    yake = support["found_by_yake"]
    tfidf = support["found_by_tfidf"]
    mpnet = support["found_by_mpnet"]

    rules = {
        "E01_yake_baseline": yake,
        "E02_union_yake_tfidf_mpnet": yake | tfidf | mpnet,
        "E03_yake_and_tfidf": yake & tfidf,
        "E04_yake_and_mpnet": yake & mpnet,
        "E05_yake_and_tfidf_or_mpnet": yake & (tfidf | mpnet),
        "E06_majority_any_2_of_3": support["vote_count"] >= 2,
        "E07_all_3_intersection": yake & tfidf & mpnet,
    }

    ensembles = {}
    for name, mask in rules.items():
        frame = support[mask].copy()
        frame.insert(0, "ensemble", name)
        frame = frame.sort_values(
            ["vote_count", "best_source_rank", "average_source_rank"],
            ascending=[False, True, True],
        ).reset_index(drop=True)
        frame["ensemble_rank"] = range(1, len(frame) + 1)
        ensembles[name] = frame

    return ensembles


def evaluate_ensemble(
    ensemble_name: str,
    candidates: pd.DataFrame,
    strict_gt: pd.DataFrame,
    review_gt: pd.DataFrame,
    negatives: pd.DataFrame,
) -> tuple[dict, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    strict_lookup = strict_gt.drop_duplicates("normalized_gt").set_index(
        "normalized_gt"
    )
    review_lookup = review_gt.drop_duplicates("normalized_gt").set_index(
        "normalized_gt"
    )
    negative_lookup = negatives.drop_duplicates("normalized_negative").set_index(
        "normalized_negative"
    )

    details = candidates.copy()
    details["match_type"] = details["normalized_candidate"].map(
        lambda term: (
            "strict_exact"
            if term in strict_lookup.index
            else (
                "manual_review_exact"
                if term in review_lookup.index
                else "unmatched"
            )
        )
    )
    details["evaluation_flag"] = details["match_type"].map(
        {
            "strict_exact": "TP",
            "manual_review_exact": "REVIEW",
            "unmatched": "FP",
        }
    )
    details["matched_gt_id"] = details["normalized_candidate"].map(
        lambda term: (
            strict_lookup.at[term, "gt_id"]
            if term in strict_lookup.index
            else (
                review_lookup.at[term, "gt_id"]
                if term in review_lookup.index
                else ""
            )
        )
    )
    details["gt_category"] = details["normalized_candidate"].map(
        lambda term: (
            strict_lookup.at[term, "gt_category"]
            if term in strict_lookup.index
            else (
                review_lookup.at[term, "gt_category"]
                if term in review_lookup.index
                else ""
            )
        )
    )
    details["negative_benchmark_hit"] = details["normalized_candidate"].isin(
        negative_lookup.index
    )

    # Manual-review matches are reported separately and excluded from strict scoring.
    scored = details[details["match_type"].ne("manual_review_exact")]
    predicted = set(scored["normalized_candidate"])
    gold = set(strict_gt["normalized_gt"])

    tp = len(predicted & gold)
    fp = len(predicted - gold)
    fn = len(gold - predicted)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0

    weights = (
        strict_gt["priority"]
        .fillna("low")
        .astype(str)
        .str.casefold()
        .map(PRIORITY_WEIGHTS)
        .fillna(1.0)
    )
    matched_mask = strict_gt["normalized_gt"].isin(predicted)
    weighted_recall = (
        float(weights[matched_mask].sum() / weights.sum())
        if weights.sum()
        else 0.0
    )

    metric = {
        "ensemble": ensemble_name,
        "candidate_count": len(details),
        "strict_scored_prediction_count": len(scored),
        "strict_gold_count": len(gold),
        "true_positives": tp,
        "false_positives": fp,
        "false_negatives": fn,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "priority_weighted_recall": weighted_recall,
        "manual_review_matches": int(
            details["match_type"].eq("manual_review_exact").sum()
        ),
        "negative_term_hits": int(details["negative_benchmark_hit"].sum()),
    }

    false_negatives = strict_gt[
        ~strict_gt["normalized_gt"].isin(predicted)
    ].copy()
    false_negatives.insert(0, "ensemble", ensemble_name)
    false_negatives["evaluation_flag"] = "FN"

    category_rows = []
    for category, group in strict_gt.groupby("gt_category"):
        category_gold = set(group["normalized_gt"])
        matched = category_gold & predicted
        category_rows.append(
            {
                "ensemble": ensemble_name,
                "gt_category": category,
                "gold_count": len(category_gold),
                "matched_count": len(matched),
                "missed_count": len(category_gold) - len(matched),
                "recall": len(matched) / len(category_gold) if category_gold else 0.0,
            }
        )

    return metric, details, false_negatives, pd.DataFrame(category_rows)


def format_excel(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        if worksheet.max_row > 0 and worksheet.max_column > 0:
            worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        for column in range(1, worksheet.max_column + 1):
            letter = get_column_letter(column)
            lengths = [
                len(str(worksheet.cell(row=row, column=column).value or ""))
                for row in range(1, min(worksheet.max_row, 200) + 1)
            ]
            worksheet.column_dimensions[letter].width = min(
                max(max(lengths, default=10) + 2, 12),
                42,
            )

        if worksheet.title == "Metrics Summary":
            percentage_headers = {
                "precision",
                "recall",
                "f1",
                "priority_weighted_recall",
            }
            header_map = {
                worksheet.cell(row=1, column=column).value: column
                for column in range(1, worksheet.max_column + 1)
            }
            for header in percentage_headers:
                if header in header_map:
                    for row in range(2, worksheet.max_row + 1):
                        worksheet.cell(
                            row=row, column=header_map[header]
                        ).number_format = "0.00%"

    workbook.save(path)


def main() -> None:
    validate_inputs()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    by_ensemble = OUTPUT_DIR / "by_ensemble"
    by_ensemble.mkdir(parents=True, exist_ok=True)

    strict_gt, review_gt, negatives = load_ground_truth()
    method_frames, combined = load_method_candidates()
    support = create_support_table(method_frames, combined)
    ensembles = build_ensembles(support)

    metrics = []
    all_details = []
    all_false_negatives = []
    all_categories = []

    for ensemble_name, candidates in ensembles.items():
        metric, details, false_negatives, category_report = evaluate_ensemble(
            ensemble_name,
            candidates,
            strict_gt,
            review_gt,
            negatives,
        )
        metrics.append(metric)
        all_details.append(details)
        all_false_negatives.append(false_negatives)
        all_categories.append(category_report)

        ensemble_dir = by_ensemble / ensemble_name
        ensemble_dir.mkdir(parents=True, exist_ok=True)
        candidates.to_csv(ensemble_dir / "candidates.csv", index=False)
        details.to_csv(ensemble_dir / "evaluation.csv", index=False)
        false_negatives.to_csv(ensemble_dir / "false_negatives.csv", index=False)
        category_report.to_csv(ensemble_dir / "category_recall.csv", index=False)
        pd.DataFrame([metric]).to_csv(ensemble_dir / "metrics.csv", index=False)

    metrics_df = pd.DataFrame(metrics).sort_values(
        ["f1", "precision", "recall"],
        ascending=[False, False, False],
    ).reset_index(drop=True)
    details_df = pd.concat(all_details, ignore_index=True)
    false_negatives_df = pd.concat(all_false_negatives, ignore_index=True)
    categories_df = pd.concat(all_categories, ignore_index=True)

    metrics_df.to_csv(OUTPUT_DIR / "ensemble_metrics_summary.csv", index=False)
    support.to_csv(OUTPUT_DIR / "method_support_matrix.csv", index=False)
    details_df.to_csv(OUTPUT_DIR / "all_ensemble_evaluations.csv", index=False)
    false_negatives_df.to_csv(
        OUTPUT_DIR / "all_ensemble_false_negatives.csv", index=False
    )
    categories_df.to_csv(OUTPUT_DIR / "all_ensemble_category_recall.csv", index=False)

    report_file = OUTPUT_DIR / "ensemble_extraction_report.xlsx"
    readme = pd.DataFrame(
        {
            "Field": [
                "Purpose",
                "Methods",
                "Candidate policy",
                "Primary selection metric",
                "Important interpretation",
            ],
            "Value": [
                "Compare logical ensembles of YAKE, TF-IDF, and KeyBERT MPNet.",
                "YAKE; TF-IDF; KeyBERT MPNet",
                "No top-k cutoff. Ensembles operate on complete unique candidate sets.",
                "Full-set exact-match F1 against strict S1 ground truth.",
                "The best ensemble on this benchmark is a development result and should be validated on held-out data.",
            ],
        }
    )

    with pd.ExcelWriter(report_file, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        metrics_df.to_excel(writer, sheet_name="Metrics Summary", index=False)
        categories_df.to_excel(writer, sheet_name="Category Recall", index=False)
        details_df.to_excel(writer, sheet_name="Evaluations", index=False)
        false_negatives_df.to_excel(writer, sheet_name="False Negatives", index=False)
        support.to_excel(writer, sheet_name="Method Support", index=False)

    format_excel(report_file)

    manifest = {
        "stage": "03b_extraction_ensembles",
        "input_directory": str(EXTRACTION_DIR),
        "ground_truth": str(GROUND_TRUTH_FILE),
        "methods": list(METHOD_FILES),
        "ensemble_rules": list(ensembles),
        "candidate_policy": "No top-k cutoff",
        "best_ensemble_by_f1": (
            metrics_df.iloc[0]["ensemble"] if not metrics_df.empty else None
        ),
    }
    (OUTPUT_DIR / "run_manifest.json").write_text(
        json.dumps(manifest, indent=2), encoding="utf-8"
    )

    print("\n" + "=" * 80)
    print("Stage 03b ensemble evaluation completed")
    print("=" * 80)
    print(metrics_df.to_string(index=False))
    print(f"\nCSV summary: {OUTPUT_DIR / 'ensemble_metrics_summary.csv'}")
    print(f"Excel report: {report_file}")


if __name__ == "__main__":
    main()
