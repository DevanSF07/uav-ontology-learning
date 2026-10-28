from __future__ import annotations

import json
import re
import time
import traceback
from pathlib import Path

import numpy as np
import pandas as pd
from sentence_transformers import SentenceTransformer


PROJECT_DIR = Path("ontology_experiment_framework")
INPUT_FILE = PROJECT_DIR / "stage_outputs/03c_weighted_voting/best_f1_candidates.csv"
GROUND_TRUTH_FILE = PROJECT_DIR / "data/Ground_Truth.xlsx"
OUTPUT_DIR = PROJECT_DIR / "stage_outputs/05_seed_similarity_filtering"

GROUND_TRUTH_SHEET = "S1_GT_TERMS"
NEGATIVE_TERMS_SHEET = "S7_NEGATIVE_TERMS"
MANUAL_REVIEW_STATUS = "needs_manual_review"

MODELS = {
    "minilm": "all-MiniLM-L6-v2",
    "mpnet": "all-mpnet-base-v2",
}

REPRESENTATIONS = {
    "term": "term",
    "term_sentence": "term_sentence",
    "term_context": "term_context",
}

# Fixed, manually declared domain anchors. These are not generated from GT.
DOMAIN_SEEDS = [
    "UAV unmanned aerial vehicle",
    "unmanned aircraft",
    "unmanned aircraft system",
    "aircraft",
    "remotely piloted aircraft",
    "remotely piloted aircraft system",
    "aerial vehicle",
    "flight control system",
    "ground control station",
    "aircraft mission",
]

MAX_SIMILARITY_THRESHOLDS = [
    0.10, 0.15, 0.20, 0.25, 0.30, 0.35, 0.40,
    0.45, 0.50, 0.55, 0.60, 0.65, 0.70,
]

# Additional multi-seed support rules.
MULTI_SEED_RULES = [
    (2, 0.30),
    (2, 0.40),
    (2, 0.50),
    (3, 0.30),
    (3, 0.40),
    (3, 0.50),
]

PRIORITY_WEIGHTS = {
    "high": 3.0,
    "medium": 2.0,
    "low": 1.0,
}


def normalize_term(value: object) -> str:
    if value is None or pd.isna(value):
        return ""

    text = str(value).casefold().strip()
    for dash in ("‐", "‑", "‒", "–", "—"):
        text = text.replace(dash, "-")

    text = re.sub(r"(?<=\w)-(?=\w)", " ", text)
    text = re.sub(r"[^a-z0-9&/+ ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def safe_text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def validate_inputs() -> None:
    missing = []

    if not INPUT_FILE.is_file():
        missing.append(str(INPUT_FILE))

    if not GROUND_TRUTH_FILE.is_file():
        missing.append(str(GROUND_TRUTH_FILE))

    if missing:
        raise FileNotFoundError(
            "Missing required files:\n" + "\n".join(missing)
        )


def load_candidates() -> pd.DataFrame:
    candidates = pd.read_csv(INPUT_FILE, encoding="utf-8")

    required = {
        "candidate",
        "normalized_candidate",
    }
    missing = required - set(candidates.columns)

    if missing:
        raise ValueError(
            f"{INPUT_FILE} is missing columns: {sorted(missing)}"
        )

    for optional_column in ("source_sentence", "context_text"):
        if optional_column not in candidates.columns:
            candidates[optional_column] = ""

    candidates = candidates.copy()
    candidates["candidate"] = candidates["candidate"].map(safe_text)
    candidates["normalized_candidate"] = candidates[
        "normalized_candidate"
    ].map(normalize_term)
    candidates["source_sentence"] = candidates["source_sentence"].map(
        safe_text
    )
    candidates["context_text"] = candidates["context_text"].map(safe_text)

    candidates = candidates[
        candidates["normalized_candidate"].ne("")
    ].copy()
    candidates = candidates.drop_duplicates(
        "normalized_candidate",
        keep="first",
    ).reset_index(drop=True)

    if candidates.empty:
        raise ValueError("The input candidate inventory is empty.")

    return candidates


def load_ground_truth():
    gt = pd.read_excel(
        GROUND_TRUTH_FILE,
        sheet_name=GROUND_TRUTH_SHEET,
        engine="openpyxl",
    )
    negatives = pd.read_excel(
        GROUND_TRUTH_FILE,
        sheet_name=NEGATIVE_TERMS_SHEET,
        engine="openpyxl",
    )

    gt = gt[
        gt["keep_for_tasks"]
        .astype(str)
        .str.contains("term_extraction", case=False, na=False)
    ].copy()
    gt["normalized_gt"] = gt["term_canonical"].map(normalize_term)
    gt = gt[gt["normalized_gt"].ne("")].copy()

    status = (
        gt["grounded_status_v2"]
        .fillna("")
        .astype(str)
        .str.casefold()
    )
    strict_gt = gt[status.ne(MANUAL_REVIEW_STATUS)].copy()
    review_gt = gt[status.eq(MANUAL_REVIEW_STATUS)].copy()

    negatives["normalized_negative"] = negatives["candidate"].map(
        normalize_term
    )
    negatives = negatives[
        negatives["normalized_negative"].ne("")
    ].copy()

    return strict_gt, review_gt, negatives


def build_representations(
    candidates: pd.DataFrame,
    representation_name: str,
) -> list[str]:
    representations = []

    for row in candidates.itertuples():
        term = safe_text(row.candidate)
        sentence = safe_text(row.source_sentence)
        context = safe_text(row.context_text)

        if representation_name == "term":
            value = term
        elif representation_name == "term_sentence":
            value = (
                f"{term}. Context sentence: {sentence}"
                if sentence
                else term
            )
        elif representation_name == "term_context":
            value = (
                f"{term}. Context: {context}"
                if context
                else term
            )
        else:
            raise ValueError(
                f"Unknown representation: {representation_name}"
            )

        representations.append(value)

    return representations


def score_candidates(
    candidates: pd.DataFrame,
    model: SentenceTransformer,
    model_key: str,
    model_name: str,
    representation_name: str,
) -> pd.DataFrame:
    representations = build_representations(
        candidates,
        representation_name,
    )

    candidate_embeddings = model.encode(
        representations,
        normalize_embeddings=True,
        show_progress_bar=True,
        batch_size=32,
    )
    seed_embeddings = model.encode(
        DOMAIN_SEEDS,
        normalize_embeddings=True,
        show_progress_bar=False,
        batch_size=32,
    )

    # Normalized embeddings make dot product equal cosine similarity.
    similarity_matrix = np.asarray(candidate_embeddings) @ np.asarray(
        seed_embeddings
    ).T

    result = candidates.copy()
    result["model_key"] = model_key
    result["model_name"] = model_name
    result["representation"] = representation_name
    result["embedding_input"] = representations

    for seed_index, seed in enumerate(DOMAIN_SEEDS):
        result[f"seed_{seed_index + 1:02d}_similarity"] = (
            similarity_matrix[:, seed_index]
        )

    closest_indices = similarity_matrix.argmax(axis=1)
    result["max_seed_similarity"] = similarity_matrix.max(axis=1)
    result["mean_seed_similarity"] = similarity_matrix.mean(axis=1)
    result["median_seed_similarity"] = np.median(
        similarity_matrix,
        axis=1,
    )
    result["closest_seed"] = [
        DOMAIN_SEEDS[index]
        for index in closest_indices
    ]

    for threshold in (0.30, 0.40, 0.50):
        result[f"seed_count_ge_{threshold:.2f}"] = (
            similarity_matrix >= threshold
        ).sum(axis=1)

    return result


def annotate_matches(
    candidates: pd.DataFrame,
    strict_gt: pd.DataFrame,
    review_gt: pd.DataFrame,
    negatives: pd.DataFrame,
) -> pd.DataFrame:
    strict_lookup = (
        strict_gt
        .drop_duplicates("normalized_gt")
        .set_index("normalized_gt")
    )
    review_lookup = (
        review_gt
        .drop_duplicates("normalized_gt")
        .set_index("normalized_gt")
    )
    negative_set = set(
        negatives["normalized_negative"].dropna()
    )

    result = candidates.copy()
    result["match_type"] = result["normalized_candidate"].map(
        lambda term: (
            "strict_exact"
            if term in strict_lookup.index
            else (
                "manual_review_exact"
                if term in review_lookup.index
                else "unmatched"
            )
        )
    )
    result["evaluation_flag_before_filter"] = result["match_type"].map(
        {
            "strict_exact": "TP",
            "manual_review_exact": "REVIEW",
            "unmatched": "FP",
        }
    )
    result["matched_gt_id"] = result["normalized_candidate"].map(
        lambda term: (
            strict_lookup.at[term, "gt_id"]
            if term in strict_lookup.index
            else (
                review_lookup.at[term, "gt_id"]
                if term in review_lookup.index
                else ""
            )
        )
    )
    result["gt_category"] = result["normalized_candidate"].map(
        lambda term: (
            strict_lookup.at[term, "gt_category"]
            if term in strict_lookup.index
            else (
                review_lookup.at[term, "gt_category"]
                if term in review_lookup.index
                else ""
            )
        )
    )
    result["negative_benchmark_hit"] = result[
        "normalized_candidate"
    ].isin(negative_set)

    return result


def evaluate_mask(
    experiment_name: str,
    annotated: pd.DataFrame,
    keep_mask: pd.Series,
    strict_gt: pd.DataFrame,
) -> tuple[dict, pd.DataFrame, pd.DataFrame]:
    retained = annotated[keep_mask].copy()
    dropped = annotated[~keep_mask].copy()

    strict_unique = strict_gt.drop_duplicates("normalized_gt").copy()
    gold = set(strict_unique["normalized_gt"])

    scored_retained = retained[
        retained["match_type"].ne("manual_review_exact")
    ]
    predicted = set(scored_retained["normalized_candidate"])

    tp = len(predicted & gold)
    fp = len(predicted - gold)
    fn = len(gold - predicted)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = (
        2 * precision * recall / (precision + recall)
        if precision + recall
        else 0.0
    )

    gt_weights = (
        strict_unique["priority"]
        .fillna("low")
        .astype(str)
        .str.casefold()
        .map(PRIORITY_WEIGHTS)
        .fillna(1.0)
    )
    matched_mask = strict_unique["normalized_gt"].isin(predicted)
    weighted_recall = (
        float(gt_weights[matched_mask].sum() / gt_weights.sum())
        if gt_weights.sum()
        else 0.0
    )

    metric = {
        "experiment": experiment_name,
        "model_key": annotated["model_key"].iloc[0],
        "model_name": annotated["model_name"].iloc[0],
        "representation": annotated["representation"].iloc[0],
        "input_candidate_count": len(annotated),
        "retained_candidate_count": len(retained),
        "dropped_candidate_count": len(dropped),
        "true_positives": tp,
        "false_positives": fp,
        "false_negatives": fn,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "priority_weighted_recall": weighted_recall,
        "tp_retained": int(
            retained["match_type"].eq("strict_exact").sum()
        ),
        "tp_dropped": int(
            dropped["match_type"].eq("strict_exact").sum()
        ),
        "fp_retained": int(
            retained["match_type"].eq("unmatched").sum()
        ),
        "fp_dropped": int(
            dropped["match_type"].eq("unmatched").sum()
        ),
        "manual_review_retained": int(
            retained["match_type"].eq("manual_review_exact").sum()
        ),
        "negative_hits_retained": int(
            retained["negative_benchmark_hit"].sum()
        ),
        "negative_hits_dropped": int(
            dropped["negative_benchmark_hit"].sum()
        ),
    }

    retained["filter_experiment"] = experiment_name
    retained["filter_decision"] = "retained"
    dropped["filter_experiment"] = experiment_name
    dropped["filter_decision"] = "dropped"

    return metric, retained, dropped


def build_filter_experiments(
    annotated: pd.DataFrame,
) -> dict[str, pd.Series]:
    experiments = {
        "baseline_no_filter": pd.Series(
            True,
            index=annotated.index,
        )
    }

    for threshold in MAX_SIMILARITY_THRESHOLDS:
        name = f"max_similarity_ge_{threshold:.2f}"
        experiments[name] = (
            annotated["max_seed_similarity"] >= threshold
        )

    for required_seed_count, similarity_threshold in MULTI_SEED_RULES:
        column = f"seed_count_ge_{similarity_threshold:.2f}"
        name = (
            f"at_least_{required_seed_count}_seeds_ge_"
            f"{similarity_threshold:.2f}"
        )
        experiments[name] = (
            annotated[column] >= required_seed_count
        )

    return experiments


def format_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False

        if worksheet.max_row > 0 and worksheet.max_column > 0:
            worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )
            cell.font = Font(
                color="FFFFFF",
                bold=True,
            )
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="center",
            )

        for column in range(1, worksheet.max_column + 1):
            lengths = [
                len(str(worksheet.cell(row, column).value or ""))
                for row in range(
                    1,
                    min(worksheet.max_row, 150) + 1,
                )
            ]
            worksheet.column_dimensions[
                get_column_letter(column)
            ].width = min(
                max(max(lengths, default=10) + 2, 11),
                42,
            )

        if worksheet.title == "Metrics Summary":
            headers = {
                worksheet.cell(1, column).value: column
                for column in range(1, worksheet.max_column + 1)
            }
            for header in (
                "precision",
                "recall",
                "f1",
                "priority_weighted_recall",
            ):
                if header in headers:
                    for row in range(2, worksheet.max_row + 1):
                        worksheet.cell(
                            row,
                            headers[header],
                        ).number_format = "0.00%"

    workbook.save(path)


def main() -> None:
    validate_inputs()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    scored_dir = OUTPUT_DIR / "scored_candidates"
    filters_dir = OUTPUT_DIR / "by_filter"
    scored_dir.mkdir(parents=True, exist_ok=True)
    filters_dir.mkdir(parents=True, exist_ok=True)

    candidates = load_candidates()
    strict_gt, review_gt, negatives = load_ground_truth()

    all_metrics = []
    all_best_retained = []
    manifest = {
        "stage": "05_seed_similarity_filtering",
        "input_file": str(INPUT_FILE),
        "ground_truth_file": str(GROUND_TRUTH_FILE),
        "input_candidate_count": len(candidates),
        "domain_seeds": DOMAIN_SEEDS,
        "models": MODELS,
        "representations": list(REPRESENTATIONS),
        "max_similarity_thresholds": MAX_SIMILARITY_THRESHOLDS,
        "multi_seed_rules": MULTI_SEED_RULES,
        "runs": [],
    }

    for model_key, model_name in MODELS.items():
        print(f"\nLoading model: {model_name}")
        model = SentenceTransformer(model_name)

        for representation_name in REPRESENTATIONS:
            run_name = f"{model_key}__{representation_name}"
            print("\n" + "=" * 80)
            print(f"Scoring: {run_name}")
            print("=" * 80)
            started = time.time()

            try:
                scored = score_candidates(
                    candidates=candidates,
                    model=model,
                    model_key=model_key,
                    model_name=model_name,
                    representation_name=representation_name,
                )
                annotated = annotate_matches(
                    scored,
                    strict_gt,
                    review_gt,
                    negatives,
                )
                annotated.to_csv(
                    scored_dir / f"{run_name}.csv",
                    index=False,
                    encoding="utf-8",
                )

                run_metrics = []
                run_outputs = []
                experiments = build_filter_experiments(annotated)

                for experiment_name, keep_mask in experiments.items():
                    full_experiment_name = (
                        f"{run_name}__{experiment_name}"
                    )
                    metric, retained, dropped = evaluate_mask(
                        full_experiment_name,
                        annotated,
                        keep_mask,
                        strict_gt,
                    )
                    run_metrics.append(metric)
                    all_metrics.append(metric)
                    run_outputs.append(
                        (metric, retained, dropped)
                    )

                run_metrics_df = pd.DataFrame(run_metrics).sort_values(
                    ["f1", "precision", "recall"],
                    ascending=[False, False, False],
                ).reset_index(drop=True)
                run_metrics_df.to_csv(
                    filters_dir / f"{run_name}__metrics.csv",
                    index=False,
                    encoding="utf-8",
                )

                best_experiment = run_metrics_df.iloc[0]["experiment"]
                best_tuple = next(
                    item
                    for item in run_outputs
                    if item[0]["experiment"] == best_experiment
                )
                _, best_retained, best_dropped = best_tuple
                best_retained.to_csv(
                    filters_dir / f"{run_name}__best_retained.csv",
                    index=False,
                    encoding="utf-8",
                )
                best_dropped.to_csv(
                    filters_dir / f"{run_name}__best_dropped.csv",
                    index=False,
                    encoding="utf-8",
                )
                all_best_retained.append(best_retained)

                manifest["runs"].append(
                    {
                        "run": run_name,
                        "status": "completed",
                        "duration_seconds": round(
                            time.time() - started,
                            3,
                        ),
                        "best_experiment": best_experiment,
                        "best_f1": float(
                            run_metrics_df.iloc[0]["f1"]
                        ),
                    }
                )

                print(
                    run_metrics_df[
                        [
                            "experiment",
                            "retained_candidate_count",
                            "true_positives",
                            "false_positives",
                            "precision",
                            "recall",
                            "f1",
                            "tp_dropped",
                            "fp_dropped",
                        ]
                    ].head(8).to_string(index=False)
                )

            except Exception as error:
                error_file = OUTPUT_DIR / f"ERROR_{run_name}.txt"
                error_file.write_text(
                    traceback.format_exc(),
                    encoding="utf-8",
                )
                manifest["runs"].append(
                    {
                        "run": run_name,
                        "status": "failed",
                        "duration_seconds": round(
                            time.time() - started,
                            3,
                        ),
                        "error": str(error),
                        "error_file": str(error_file),
                    }
                )
                print(f"Failed {run_name}: {error}")

        del model

    metrics_df = pd.DataFrame(all_metrics).sort_values(
        ["f1", "precision", "recall"],
        ascending=[False, False, False],
    ).reset_index(drop=True)
    metrics_df.to_csv(
        OUTPUT_DIR / "seed_similarity_metrics.csv",
        index=False,
        encoding="utf-8",
    )

    if metrics_df.empty:
        raise RuntimeError(
            "No seed-similarity experiment completed successfully."
        )

    best_experiment = metrics_df.iloc[0]["experiment"]
    best_model_key = metrics_df.iloc[0]["model_key"]
    best_representation = metrics_df.iloc[0]["representation"]

    matching_best_file = (
        filters_dir
        / f"{best_model_key}__{best_representation}__best_retained.csv"
    )
    best_candidates = pd.read_csv(
        matching_best_file,
        encoding="utf-8",
    )
    best_candidates.to_csv(
        OUTPUT_DIR / "best_seed_filtered_candidates.csv",
        index=False,
        encoding="utf-8",
    )

    seeds_df = pd.DataFrame(
        {
            "seed_id": range(1, len(DOMAIN_SEEDS) + 1),
            "seed_text": DOMAIN_SEEDS,
        }
    )

    report_file = OUTPUT_DIR / "seed_similarity_report.xlsx"
    readme = pd.DataFrame(
        {
            "Field": [
                "Purpose",
                "Input",
                "Candidate policy",
                "Similarity method",
                "Models",
                "Representations",
                "Best experiment",
                "Important caution",
            ],
            "Value": [
                "Filter extracted terms by semantic similarity to fixed UAV-domain seeds.",
                str(INPUT_FILE),
                "No top-k cutoff; thresholds are evaluated against the complete input inventory.",
                "Cosine similarity via dot product of normalized embeddings.",
                "; ".join(MODELS.values()),
                "; ".join(REPRESENTATIONS),
                best_experiment,
                "Threshold selection on the same GT is a development result and requires held-out validation.",
            ],
        }
    )

    with pd.ExcelWriter(
        report_file,
        engine="openpyxl",
    ) as writer:
        readme.to_excel(
            writer,
            sheet_name="README",
            index=False,
        )
        seeds_df.to_excel(
            writer,
            sheet_name="Domain Seeds",
            index=False,
        )
        metrics_df.to_excel(
            writer,
            sheet_name="Metrics Summary",
            index=False,
        )
        best_candidates.to_excel(
            writer,
            sheet_name="Best Retained Terms",
            index=False,
        )

    format_workbook(report_file)

    manifest["best_experiment_by_f1"] = best_experiment
    manifest["best_f1"] = float(metrics_df.iloc[0]["f1"])
    manifest["best_candidate_file"] = str(
        OUTPUT_DIR / "best_seed_filtered_candidates.csv"
    )
    (OUTPUT_DIR / "run_manifest.json").write_text(
        json.dumps(
            manifest,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    print("\n" + "=" * 80)
    print("Stage 05 seed-similarity filtering completed")
    print("=" * 80)
    print(
        metrics_df[
            [
                "experiment",
                "retained_candidate_count",
                "true_positives",
                "false_positives",
                "false_negatives",
                "precision",
                "recall",
                "f1",
                "tp_dropped",
                "fp_dropped",
            ]
        ].head(20).to_string(index=False)
    )
    print(f"\nBest experiment: {best_experiment}")
    print(
        "Best retained candidates: "
        f"{OUTPUT_DIR / 'best_seed_filtered_candidates.csv'}"
    )
    print(f"Excel report: {report_file}")


if __name__ == "__main__":
    main()
